"""问题一：基于储能电量状态的高精度动态规划调度。

优化部分只使用动态规划，不调用线性规划或混合整数规划求解器。
附件时间戳按区间右端点解释；数组顺序保持不变，输出标签改为物理区间。
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import time
from html import escape
from pathlib import Path
from typing import Sequence

import numpy as np


@dataclass(frozen=True)
class DispatchConfig:
    e_min: float = 1200.0
    e_max: float = 10800.0
    initial_energy: float = 6000.0
    terminal_energy: float = 6000.0
    power_max_kw: float = 5000.0
    efficiency: float = 0.9
    dt_hours: float = 1.0 / 6.0
    coarse_step_kwh: float = 5.0
    refinement_widths_kwh: tuple[float, ...] = (30.0, 3.0)
    refinement_steps_kwh: tuple[float, ...] = (0.25, 0.05)
    tolerance: float = 1e-7
    curtailment_tie_breaker: float = 1e-9

    def __post_init__(self) -> None:
        if not 0.0 < self.efficiency <= 1.0:
            raise ValueError("充放电效率必须位于 (0, 1]。")
        if not self.e_min <= self.initial_energy <= self.e_max:
            raise ValueError("初始储能量不在允许区间内。")
        if not self.e_min <= self.terminal_energy <= self.e_max:
            raise ValueError("终止储能量不在允许区间内。")
        if self.power_max_kw <= 0 or self.dt_hours <= 0:
            raise ValueError("功率上限和时段长度必须为正。")
        if self.coarse_step_kwh <= 0:
            raise ValueError("粗网格步长必须为正。")
        if len(self.refinement_widths_kwh) != len(self.refinement_steps_kwh):
            raise ValueError("局部细化宽度和步长的层数必须一致。")
        if self.curtailment_tie_breaker < 0:
            raise ValueError("弃光并列决策系数不得为负。")

    @property
    def bus_energy_limit_kwh(self) -> float:
        return self.power_max_kw * self.dt_hours

    @property
    def max_charge_soc_delta_kwh(self) -> float:
        return self.efficiency * self.bus_energy_limit_kwh

    @property
    def max_discharge_soc_delta_kwh(self) -> float:
        return self.bus_energy_limit_kwh / self.efficiency


@dataclass(frozen=True)
class DispatchSolution:
    soc_kwh: np.ndarray
    grid_kwh: np.ndarray
    charge_kwh: np.ndarray
    discharge_kwh: np.ndarray
    curtailment_kwh: np.ndarray
    cost_yuan: np.ndarray
    net_load_kwh: np.ndarray
    action_soc_delta_kwh: np.ndarray
    resolution_kwh: float


def _endpoint_minutes(value: object, index: int) -> int:
    if isinstance(value, time):
        minutes = value.hour * 60 + value.minute
        return 1440 if index == 143 and minutes == 0 else minutes
    text = str(value).strip()
    if "+1" in text:
        return 1440
    text = text.replace("：", ":")
    parts = text.split(":")
    if len(parts) < 2:
        raise ValueError(f"无法识别时间值：{value!r}")
    return int(parts[0]) * 60 + int(parts[1])


def _clock_label(minutes: int) -> str:
    if minutes == 1440:
        return "24:00"
    return f"{minutes // 60:02d}:{minutes % 60:02d}"


def build_interval_labels(raw_endpoints: Sequence[object]) -> list[str]:
    """验证 144 个右端点并返回 00:00-00:10 ... 23:50-24:00。"""
    if len(raw_endpoints) != 144:
        raise ValueError(f"时间行数应为 144，实际为 {len(raw_endpoints)}。")
    actual = [_endpoint_minutes(value, i) for i, value in enumerate(raw_endpoints)]
    expected = [10 * (i + 1) for i in range(144)]
    if actual != expected:
        mismatches = [
            (i + 1, actual[i], expected[i])
            for i in range(144)
            if actual[i] != expected[i]
        ][:5]
        raise ValueError(f"附件时间端点不连续或发生错位，前几个差异：{mismatches}")
    return [
        f"{_clock_label(10 * i)}-{_clock_label(10 * (i + 1))}"
        for i in range(144)
    ]


def read_problem1_data(path: str | Path) -> dict[str, object]:
    """按列位置读取附件1，并严格验证 144 个右端点时间戳。"""
    from openpyxl import load_workbook

    source = Path(path)
    if not source.exists():
        raise FileNotFoundError(f"找不到问题一数据文件：{source}")
    workbook = load_workbook(source, data_only=True, read_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
    finally:
        workbook.close()
    rows = [row for row in rows if row[0] is not None]
    if len(rows) != 144:
        raise ValueError(f"附件1应包含 144 个时段，实际为 {len(rows)}。")
    raw_time = [row[0] for row in rows]
    price = np.asarray([row[1] for row in rows], dtype=float)
    load_kw = np.asarray([row[2] for row in rows], dtype=float)
    pv_kw = np.asarray([row[3] for row in rows], dtype=float)
    _validate_inputs(load_kw, pv_kw, price)
    return {
        "intervals": build_interval_labels(raw_time),
        "raw_time_endpoints": raw_time,
        "price": price,
        "load_kw": load_kw,
        "pv_kw": pv_kw,
    }


def storage_bus_exchange(x_kwh: np.ndarray, efficiency: float) -> np.ndarray:
    """库存变化映射为母线侧净取电量 ψ(x)，正值取电、负值供电。"""
    x = np.asarray(x_kwh, dtype=float)
    return np.where(x >= 0.0, x / efficiency, efficiency * x)


def _validate_inputs(
    load_kw: np.ndarray, pv_kw: np.ndarray, price: np.ndarray
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    load = np.asarray(load_kw, dtype=float)
    pv = np.asarray(pv_kw, dtype=float)
    tariff = np.asarray(price, dtype=float)
    if load.ndim != 1 or pv.ndim != 1 or tariff.ndim != 1:
        raise ValueError("负荷、光伏和电价必须是一维数组。")
    if not (len(load) == len(pv) == len(tariff)) or len(load) == 0:
        raise ValueError("负荷、光伏和电价长度必须相同且非空。")
    if not np.all(np.isfinite(np.concatenate([load, pv, tariff]))):
        raise ValueError("输入数据含 NaN 或 Inf。")
    if np.any(load < 0) or np.any(pv < 0) or np.any(tariff < 0):
        raise ValueError("负荷、光伏和电价不得为负。")
    return load, pv, tariff


def _uniform_grid(config: DispatchConfig) -> np.ndarray:
    count = int(np.floor((config.e_max - config.e_min) / config.coarse_step_kwh))
    grid = config.e_min + config.coarse_step_kwh * np.arange(count + 1)
    if grid[-1] < config.e_max - config.tolerance:
        grid = np.append(grid, config.e_max)
    grid = np.unique(
        np.append(grid, [config.initial_energy, config.terminal_energy])
    )
    return np.sort(grid)


def _backward_dp_uniform(
    net_kwh: np.ndarray,
    pv_kwh: np.ndarray,
    price: np.ndarray,
    states: np.ndarray,
    config: DispatchConfig,
) -> np.ndarray:
    terminal_index = int(np.argmin(np.abs(states - config.terminal_energy)))
    if abs(states[terminal_index] - config.terminal_energy) > config.tolerance:
        raise RuntimeError("终止储能量未进入 DP 网格。")
    future = np.full(len(states), np.inf)
    future[terminal_index] = 0.0
    policies = np.full((len(net_kwh), len(states)), -1, dtype=np.int32)

    for t in range(len(net_kwh) - 1, -1, -1):
        current_value = np.full(len(states), np.inf)
        for i, current in enumerate(states):
            low = current - config.max_discharge_soc_delta_kwh
            high = current + config.max_charge_soc_delta_kwh
            j0 = int(np.searchsorted(states, low - config.tolerance, side="left"))
            j1 = int(np.searchsorted(states, high + config.tolerance, side="right"))
            if j0 >= j1:
                continue
            delta = states[j0:j1] - current
            pre_grid = net_kwh[t] + storage_bus_exchange(delta, config.efficiency)
            grid = np.maximum(pre_grid, 0.0)
            curtailment = np.maximum(-pre_grid, 0.0)
            physically_feasible = curtailment <= pv_kwh[t] + config.tolerance
            candidates = (
                price[t] * grid
                + config.curtailment_tie_breaker * curtailment
                + future[j0:j1]
            )
            candidates[~physically_feasible] = np.inf
            local = int(np.argmin(candidates))
            if np.isfinite(candidates[local]):
                current_value[i] = candidates[local]
                policies[t, i] = j0 + local
        future = current_value

    start_index = int(np.argmin(np.abs(states - config.initial_energy)))
    if not np.isfinite(future[start_index]):
        raise RuntimeError("给定首尾 SOC 和功率上限下不存在可行策略。")
    path = np.empty(len(net_kwh) + 1)
    path[0] = states[start_index]
    index = start_index
    for t in range(len(net_kwh)):
        index = int(policies[t, index])
        if index < 0:
            raise RuntimeError(f"DP 回溯在第 {t + 1} 个时段失败。")
        path[t + 1] = states[index]
    return path


def _local_state_grid(
    center: float, width: float, step: float, config: DispatchConfig
) -> np.ndarray:
    low = max(config.e_min, center - width)
    high = min(config.e_max, center + width)
    count = int(np.floor((high - low) / step))
    values = low + step * np.arange(count + 1)
    if values[-1] < high - config.tolerance:
        values = np.append(values, high)
    return np.unique(np.round(np.append(values, center), 10))


def _backward_dp_variable_grids(
    net_kwh: np.ndarray,
    pv_kwh: np.ndarray,
    price: np.ndarray,
    centers: np.ndarray,
    width: float,
    step: float,
    config: DispatchConfig,
) -> np.ndarray:
    state_grids: list[np.ndarray] = [np.array([config.initial_energy])]
    state_grids.extend(
        _local_state_grid(center, width, step, config)
        for center in centers[1:-1]
    )
    state_grids.append(np.array([config.terminal_energy]))

    future = np.array([0.0])
    policies: list[np.ndarray] = [np.empty(0, dtype=np.int32)] * len(net_kwh)
    for t in range(len(net_kwh) - 1, -1, -1):
        current = state_grids[t]
        nxt = state_grids[t + 1]
        delta = nxt[None, :] - current[:, None]
        feasible = (
            (delta >= -config.max_discharge_soc_delta_kwh - config.tolerance)
            & (delta <= config.max_charge_soc_delta_kwh + config.tolerance)
        )
        pre_grid = net_kwh[t] + storage_bus_exchange(delta, config.efficiency)
        grid = np.maximum(pre_grid, 0.0)
        curtailment = np.maximum(-pre_grid, 0.0)
        feasible &= curtailment <= pv_kwh[t] + config.tolerance
        candidates = (
            price[t] * grid
            + config.curtailment_tie_breaker * curtailment
            + future[None, :]
        )
        candidates[~feasible] = np.inf
        choice = np.argmin(candidates, axis=1).astype(np.int32)
        value = candidates[np.arange(len(current)), choice]
        choice[~np.isfinite(value)] = -1
        policies[t] = choice
        future = value

    if not np.isfinite(future[0]):
        raise RuntimeError(
            f"局部细化网格不可行（宽度 {width} kWh，步长 {step} kWh）。"
        )
    path = np.empty(len(net_kwh) + 1)
    path[0] = config.initial_energy
    index = 0
    for t in range(len(net_kwh)):
        index = int(policies[t][index])
        if index < 0:
            raise RuntimeError(f"局部 DP 回溯在第 {t + 1} 个时段失败。")
        path[t + 1] = state_grids[t + 1][index]
    return path


def _solution_from_soc(
    load_kw: np.ndarray,
    pv_kw: np.ndarray,
    price: np.ndarray,
    soc_kwh: np.ndarray,
    resolution_kwh: float,
    config: DispatchConfig,
) -> DispatchSolution:
    net = (load_kw - pv_kw) * config.dt_hours
    delta = np.diff(soc_kwh)
    charge = np.where(delta > config.tolerance, delta / config.efficiency, 0.0)
    discharge = np.where(
        delta < -config.tolerance, -config.efficiency * delta, 0.0
    )
    pre_grid = net + charge - discharge
    grid = np.maximum(pre_grid, 0.0)
    curtailment = np.maximum(-pre_grid, 0.0)
    cost = price * grid
    return DispatchSolution(
        soc_kwh=soc_kwh,
        grid_kwh=grid,
        charge_kwh=charge,
        discharge_kwh=discharge,
        curtailment_kwh=curtailment,
        cost_yuan=cost,
        net_load_kwh=net,
        action_soc_delta_kwh=delta,
        resolution_kwh=resolution_kwh,
    )


def run_adaptive_dp(
    load_kw: np.ndarray,
    pv_kw: np.ndarray,
    price: np.ndarray,
    config: DispatchConfig | None = None,
) -> DispatchSolution:
    """先做全局 DP，再围绕可行路径逐级细化状态网格。"""
    config = config or DispatchConfig()
    load, pv, tariff = _validate_inputs(load_kw, pv_kw, price)
    net = (load - pv) * config.dt_hours
    pv_energy = pv * config.dt_hours
    path = _backward_dp_uniform(
        net, pv_energy, tariff, _uniform_grid(config), config
    )
    resolution = config.coarse_step_kwh
    for width, step in zip(
        config.refinement_widths_kwh, config.refinement_steps_kwh
    ):
        if width <= 0 or step <= 0:
            raise ValueError("细化宽度和步长必须为正。")
        path = _backward_dp_variable_grids(
            net, pv_energy, tariff, path, width, step, config
        )
        resolution = step
    return _solution_from_soc(load, pv, tariff, path, resolution, config)


def validate_dispatch(
    load_kw: np.ndarray,
    pv_kw: np.ndarray,
    price: np.ndarray,
    soc_kwh: np.ndarray,
    grid_kwh: np.ndarray,
    charge_kwh: np.ndarray,
    discharge_kwh: np.ndarray,
    curtailment_kwh: np.ndarray,
    config: DispatchConfig | None = None,
) -> dict[str, bool | float | int]:
    config = config or DispatchConfig()
    load, pv, tariff = _validate_inputs(load_kw, pv_kw, price)
    soc = np.asarray(soc_kwh, dtype=float)
    grid = np.asarray(grid_kwh, dtype=float)
    charge = np.asarray(charge_kwh, dtype=float)
    discharge = np.asarray(discharge_kwh, dtype=float)
    curtailment = np.asarray(curtailment_kwh, dtype=float)
    t_count = len(load)
    lengths_ok = (
        len(soc) == t_count + 1
        and all(len(a) == t_count for a in [grid, charge, discharge, curtailment])
    )
    if not lengths_ok:
        return {
            "complete_144_slots": t_count == 144,
            "array_lengths_ok": False,
            "finite_values_ok": False,
            "soc_within_bounds": False,
            "charge_power_within_limit": False,
            "discharge_power_within_limit": False,
            "soc_dynamics_ok": False,
            "energy_balance_ok": False,
            "no_reverse_sale": False,
            "curtailment_nonnegative": False,
            "curtailment_within_available_pv": False,
            "initial_soc_ok": False,
            "terminal_soc_ok": False,
            "no_simultaneous_charge_discharge": False,
            "max_balance_error_kwh": float("inf"),
            "max_soc_dynamics_error_kwh": float("inf"),
            "max_curtailment_excess_kwh": float("inf"),
        }

    arrays = [load, pv, tariff, soc, grid, charge, discharge, curtailment]
    finite = all(np.all(np.isfinite(a)) for a in arrays)
    soc_dynamics_error = np.diff(soc) - (
        config.efficiency * charge - discharge / config.efficiency
    )
    balance_error = (
        grid
        + pv * config.dt_hours
        + discharge
        - load * config.dt_hours
        - charge
        - curtailment
    )
    tol = max(config.tolerance, 1e-6)
    pv_energy = pv * config.dt_hours
    return {
        "complete_144_slots": t_count == 144,
        "array_lengths_ok": lengths_ok,
        "finite_values_ok": finite,
        "soc_within_bounds": bool(
            np.min(soc) >= config.e_min - tol
            and np.max(soc) <= config.e_max + tol
        ),
        "charge_power_within_limit": bool(
            np.max(charge, initial=0.0) <= config.bus_energy_limit_kwh + tol
        ),
        "discharge_power_within_limit": bool(
            np.max(discharge, initial=0.0) <= config.bus_energy_limit_kwh + tol
        ),
        "soc_dynamics_ok": bool(np.max(np.abs(soc_dynamics_error)) <= tol),
        "energy_balance_ok": bool(np.max(np.abs(balance_error)) <= tol),
        "no_reverse_sale": bool(np.min(grid, initial=0.0) >= -tol),
        "curtailment_nonnegative": bool(np.min(curtailment, initial=0.0) >= -tol),
        "curtailment_within_available_pv": bool(
            np.max(curtailment - pv_energy, initial=0.0) <= tol
        ),
        "initial_soc_ok": bool(abs(soc[0] - config.initial_energy) <= tol),
        "terminal_soc_ok": bool(abs(soc[-1] - config.terminal_energy) <= tol),
        "no_simultaneous_charge_discharge": bool(
            not np.any((charge > tol) & (discharge > tol))
        ),
        "max_balance_error_kwh": float(np.max(np.abs(balance_error))),
        "max_soc_dynamics_error_kwh": float(
            np.max(np.abs(soc_dynamics_error))
        ),
        "max_curtailment_excess_kwh": float(
            np.max(curtailment - pv_energy, initial=0.0)
        ),
    }


def _write_svg_line_chart(
    output_path: Path,
    title: str,
    y_label: str,
    series: Sequence[tuple[str, np.ndarray, str]],
    secondary: tuple[str, np.ndarray, str, str] | None = None,
) -> None:
    """用标准库生成可缩放矢量图，避免额外绘图库依赖。"""
    width, height = 1400, 720
    left, right, top, bottom = 105, 95, 78, 88
    plot_w, plot_h = width - left - right, height - top - bottom
    count = len(series[0][1])
    primary_values = np.concatenate([np.asarray(item[1], dtype=float) for item in series])
    y_min = min(0.0, float(np.min(primary_values)))
    y_max = max(0.0, float(np.max(primary_values)))
    if abs(y_max - y_min) < 1e-12:
        y_max = y_min + 1.0
    padding = 0.06 * (y_max - y_min)
    y_min -= padding
    y_max += padding

    def x_pos(i: int) -> float:
        return left + plot_w * i / max(count - 1, 1)

    def y_pos(value: float) -> float:
        return top + plot_h * (y_max - value) / (y_max - y_min)

    elements = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        '<rect width="100%" height="100%" fill="#FFFFFF"/>',
        '<style>text{font-family:"Microsoft YaHei","SimHei",Arial,sans-serif;fill:#263238}</style>',
        f'<text x="{left}" y="38" font-size="26" font-weight="700">{escape(title)}</text>',
    ]
    for tick in range(6):
        value = y_min + (y_max - y_min) * tick / 5
        y = y_pos(value)
        elements.append(
            f'<line x1="{left}" y1="{y:.2f}" x2="{left + plot_w}" y2="{y:.2f}" stroke="#E4EAF0" stroke-width="1"/>'
        )
        elements.append(
            f'<text x="{left - 12}" y="{y + 5:.2f}" text-anchor="end" font-size="14">{value:.0f}</text>'
        )
    for hour in range(0, 25, 4):
        index = min(hour * 6, count - 1)
        x = x_pos(index)
        elements.append(
            f'<line x1="{x:.2f}" y1="{top}" x2="{x:.2f}" y2="{top + plot_h}" stroke="#F0F3F6" stroke-width="1"/>'
        )
        elements.append(
            f'<text x="{x:.2f}" y="{top + plot_h + 28}" text-anchor="middle" font-size="14">{hour:02d}:00</text>'
        )
    zero_y = y_pos(0.0)
    elements.append(
        f'<line x1="{left}" y1="{zero_y:.2f}" x2="{left + plot_w}" y2="{zero_y:.2f}" stroke="#607D8B" stroke-width="1.4"/>'
    )
    for label, values, color in series:
        points = " ".join(
            f"{x_pos(i):.2f},{y_pos(float(value)):.2f}"
            for i, value in enumerate(values)
        )
        elements.append(
            f'<polyline points="{points}" fill="none" stroke="{color}" stroke-width="2.4" stroke-linejoin="round" stroke-linecap="round"/>'
        )

    legend_x = left
    for label, _, color in series:
        elements.append(
            f'<line x1="{legend_x}" y1="62" x2="{legend_x + 34}" y2="62" stroke="{color}" stroke-width="4"/>'
        )
        elements.append(
            f'<text x="{legend_x + 43}" y="67" font-size="15">{escape(label)}</text>'
        )
        legend_x += 155

    if secondary is not None:
        sec_label, sec_values, sec_color, sec_y_label = secondary
        sec_values = np.asarray(sec_values, dtype=float)
        sec_min = float(np.min(sec_values))
        sec_max = float(np.max(sec_values))
        if abs(sec_max - sec_min) < 1e-12:
            sec_max = sec_min + 1.0
        sec_pad = 0.06 * (sec_max - sec_min)
        sec_min -= sec_pad
        sec_max += sec_pad

        def sec_y(value: float) -> float:
            return top + plot_h * (sec_max - value) / (sec_max - sec_min)

        points = " ".join(
            f"{x_pos(i):.2f},{sec_y(float(value)):.2f}"
            for i, value in enumerate(sec_values)
        )
        elements.append(
            f'<polyline points="{points}" fill="none" stroke="{sec_color}" stroke-width="2.2" stroke-dasharray="8 5"/>'
        )
        elements.append(
            f'<line x1="{legend_x}" y1="62" x2="{legend_x + 34}" y2="62" stroke="{sec_color}" stroke-width="4" stroke-dasharray="8 5"/>'
        )
        elements.append(
            f'<text x="{legend_x + 43}" y="67" font-size="15">{escape(sec_label)}</text>'
        )
        for tick in range(6):
            value = sec_min + (sec_max - sec_min) * tick / 5
            y = sec_y(value)
            elements.append(
                f'<text x="{left + plot_w + 12}" y="{y + 5:.2f}" font-size="14">{value:.2f}</text>'
            )
        elements.append(
            f'<text x="{width - 20}" y="{top + plot_h / 2}" font-size="16" text-anchor="middle" transform="rotate(90 {width - 20} {top + plot_h / 2})">{escape(sec_y_label)}</text>'
        )

    elements.extend(
        [
            f'<text x="24" y="{top + plot_h / 2}" font-size="16" text-anchor="middle" transform="rotate(-90 24 {top + plot_h / 2})">{escape(y_label)}</text>',
            f'<text x="{left + plot_w / 2}" y="{height - 24}" font-size="16" text-anchor="middle">时间</text>',
            f'<rect x="{left}" y="{top}" width="{plot_w}" height="{plot_h}" fill="none" stroke="#90A4AE" stroke-width="1.2"/>',
            '</svg>',
        ]
    )
    output_path.write_text("\n".join(elements), encoding="utf-8")


def write_plots(
    output_dir: str | Path,
    data: dict[str, object],
    solution: DispatchSolution,
    config: DispatchConfig,
) -> list[Path]:
    folder = Path(output_dir)
    folder.mkdir(parents=True, exist_ok=True)
    load = np.asarray(data["load_kw"], dtype=float)
    pv = np.asarray(data["pv_kw"], dtype=float)
    price = np.asarray(data["price"], dtype=float)
    grid_kw = solution.grid_kwh / config.dt_hours
    charge_kw = solution.charge_kwh / config.dt_hours
    discharge_kw = solution.discharge_kwh / config.dt_hours
    soc_end = solution.soc_kwh[1:]
    paths = [
        folder / "负荷_光伏_购电曲线.svg",
        folder / "储能充放电功率曲线.svg",
        folder / "储能SOC轨迹.svg",
        folder / "电价与储能行为.svg",
    ]
    _write_svg_line_chart(
        paths[0],
        "问题一：负荷、光伏与计划购电功率",
        "功率/kW",
        [
            ("负荷", load, "#263238"),
            ("光伏", pv, "#F9A825"),
            ("计划购电", grid_kw, "#1565C0"),
        ],
    )
    _write_svg_line_chart(
        paths[1],
        "问题一：储能充放电功率",
        "功率/kW（放电为负）",
        [
            ("充电功率", charge_kw, "#2E7D32"),
            ("放电功率", -discharge_kw, "#C62828"),
        ],
    )
    _write_svg_line_chart(
        paths[2],
        "问题一：储能 SOC 轨迹",
        "储能量/kWh",
        [
            ("SOC", soc_end, "#6A1B9A"),
            ("运行上限", np.full_like(soc_end, config.e_max), "#EF6C00"),
            ("运行下限", np.full_like(soc_end, config.e_min), "#00838F"),
        ],
    )
    _write_svg_line_chart(
        paths[3],
        "问题一：电价与储能行为",
        "储能功率/kW（充电为正）",
        [("净储能功率", charge_kw - discharge_kw, "#3949AB")],
        secondary=("电价", price, "#F57C00", "电价/(元/kWh)"),
    )
    return paths


def write_result_workbook(
    output_path: str | Path,
    data: dict[str, object],
    solution: DispatchSolution,
    checks: dict[str, bool | float | int],
    config: DispatchConfig,
) -> Path:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    output = Path(output_path)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "汇总"
    detail = workbook.create_sheet("DP调度结果")
    charts = workbook.create_sheet("图表")
    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False

    dark = "1F4E78"
    medium = "D9EAF7"
    light = "F4F7FA"
    green = "E2F0D9"
    red = "FCE4D6"
    thin_gray = Side(style="thin", color="D5DDE5")
    header_fill = PatternFill("solid", fgColor=dark)
    section_fill = PatternFill("solid", fgColor=medium)

    summary["A2"] = "问题一：动态规划购电与储能调度结果"
    summary["A2"].font = Font(name="Microsoft YaHei", size=15, bold=True, color="1F2937")
    summary["A3"] = "附件时间为区间右端点；数值不平移，输出按实际 10 分钟区间标注。"
    summary["A3"].font = Font(name="Microsoft YaHei", size=10, italic=True, color="52606D")
    summary["A5"] = "关键指标"
    summary["B5"] = "数值"
    summary["A5:B5"][0][0].fill = header_fill
    summary["A5:B5"][0][1].fill = header_fill
    for cell in summary[5]:
        if cell.column <= 2:
            cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")
    metrics = [
        ("全天总购电量/kWh", float(np.sum(solution.grid_kwh))),
        ("全天总购电费用/元", float(np.sum(solution.cost_yuan))),
        ("总充电量/kWh（母线侧）", float(np.sum(solution.charge_kwh))),
        ("总放电量/kWh（供负荷侧）", float(np.sum(solution.discharge_kwh))),
        ("最低SOC/kWh", float(np.min(solution.soc_kwh))),
        ("最高SOC/kWh", float(np.max(solution.soc_kwh))),
        ("0:00 SOC/kWh", float(solution.soc_kwh[0])),
        ("24:00 SOC/kWh", float(solution.soc_kwh[-1])),
        ("弃光电量/kWh", float(np.sum(solution.curtailment_kwh))),
        ("DP最终分辨率/kWh", float(solution.resolution_kwh)),
    ]
    for row_index, (label, value) in enumerate(metrics, start=6):
        summary.cell(row_index, 1, label)
        summary.cell(row_index, 2, value)
        summary.cell(row_index, 2).number_format = "#,##0.0000"

    summary["D5"] = "模型参数"
    summary["E5"] = "设定"
    for cell in summary["D5:E5"][0]:
        cell.fill = header_fill
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    parameters = [
        ("时段长度/h", config.dt_hours),
        ("储能运行下限/kWh", config.e_min),
        ("储能运行上限/kWh", config.e_max),
        ("最大充放电功率/kW", config.power_max_kw),
        ("充放电效率", config.efficiency),
        ("单时段母线侧功率上限/kWh", config.bus_energy_limit_kwh),
        ("最大充电库存增量/kWh", config.max_charge_soc_delta_kwh),
        ("最大放电库存减量/kWh", config.max_discharge_soc_delta_kwh),
        ("光伏余量处理", "允许弃光，不允许反送电"),
    ]
    for row_index, (label, value) in enumerate(parameters, start=6):
        summary.cell(row_index, 4, label)
        summary.cell(row_index, 5, value)
        if isinstance(value, (int, float)):
            summary.cell(row_index, 5).number_format = "#,##0.0000"

    summary["A18"] = "约束与完整性检查"
    summary["B18"] = "结果"
    for cell in summary["A18:B18"][0]:
        cell.fill = header_fill
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    check_labels = {
        "complete_144_slots": "144个时段完整",
        "array_lengths_ok": "数组长度一致",
        "finite_values_ok": "无NaN/Inf",
        "soc_within_bounds": "SOC上下界",
        "charge_power_within_limit": "充电功率上限",
        "discharge_power_within_limit": "放电功率上限",
        "soc_dynamics_ok": "效率与SOC递推",
        "energy_balance_ok": "逐时段电量平衡",
        "no_reverse_sale": "禁止向外网反向售电",
        "curtailment_nonnegative": "弃光量非负",
        "curtailment_within_available_pv": "弃光不超过当期光伏",
        "initial_soc_ok": "初始SOC",
        "terminal_soc_ok": "终止SOC",
        "no_simultaneous_charge_discharge": "无同时充放电",
    }
    row_index = 19
    for key, label in check_labels.items():
        passed = bool(checks[key])
        summary.cell(row_index, 1, label)
        summary.cell(row_index, 2, "通过" if passed else "失败")
        summary.cell(row_index, 2).fill = PatternFill(
            "solid", fgColor=green if passed else red
        )
        row_index += 1
    summary.cell(row_index, 1, "最大电量平衡误差/kWh")
    summary.cell(row_index, 2, float(checks["max_balance_error_kwh"]))
    summary.cell(row_index + 1, 1, "最大SOC递推误差/kWh")
    summary.cell(row_index + 1, 2, float(checks["max_soc_dynamics_error_kwh"]))
    summary.cell(row_index + 2, 1, "最大弃光超额/kWh")
    summary.cell(row_index + 2, 2, float(checks["max_curtailment_excess_kwh"]))
    summary.cell(row_index, 2).number_format = "0.0000000000E+00"
    summary.cell(row_index + 1, 2).number_format = "0.0000000000E+00"
    summary.cell(row_index + 2, 2).number_format = "0.0000000000E+00"

    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 18
    summary.column_dimensions["C"].width = 3
    summary.column_dimensions["D"].width = 31
    summary.column_dimensions["E"].width = 24
    summary.freeze_panes = "A5"
    summary.sheet_properties.tabColor = dark

    detail["A2"] = "144个10分钟时段的最优DP决策"
    detail["A2"].font = Font(name="Microsoft YaHei", size=15, bold=True, color="1F2937")
    detail["A3"] = "储能量为时段结束值；充电量和放电量均为母线侧电量。"
    detail["A3"].font = Font(name="Microsoft YaHei", size=10, italic=True, color="52606D")
    headers = [
        "序号",
        "时间",
        "负荷/kW",
        "光伏/kW",
        "净负荷/kWh",
        "电价/(元/kWh)",
        "计划购电量/kWh",
        "计划购电功率/kW",
        "充电量/kWh",
        "放电量/kWh",
        "充电功率/kW",
        "放电功率/kW",
        "储能动作ΔE/kWh",
        "储能量/kWh",
        "弃光电量/kWh",
        "单时段费用/元",
        "累计费用/元",
        "最优决策",
        "净储能功率/kW",
        "缩放电价（电价×4000）",
    ]
    for column, header in enumerate(headers, start=1):
        detail.cell(4, column, header)
    cumulative = np.cumsum(solution.cost_yuan)
    intervals = list(data["intervals"])
    load = np.asarray(data["load_kw"], dtype=float)
    pv = np.asarray(data["pv_kw"], dtype=float)
    price = np.asarray(data["price"], dtype=float)
    for i in range(144):
        if solution.charge_kwh[i] > config.tolerance:
            decision = "充电"
        elif solution.discharge_kwh[i] > config.tolerance:
            decision = "放电"
        else:
            decision = "不动作"
        detail.append(
            [
                i + 1,
                intervals[i],
                float(load[i]),
                float(pv[i]),
                float(solution.net_load_kwh[i]),
                float(price[i]),
                float(solution.grid_kwh[i]),
                float(solution.grid_kwh[i] / config.dt_hours),
                float(solution.charge_kwh[i]),
                float(solution.discharge_kwh[i]),
                float(solution.charge_kwh[i] / config.dt_hours),
                float(solution.discharge_kwh[i] / config.dt_hours),
                float(solution.action_soc_delta_kwh[i]),
                float(solution.soc_kwh[i + 1]),
                float(solution.curtailment_kwh[i]),
                float(solution.cost_yuan[i]),
                float(cumulative[i]),
                decision,
                float(
                    (solution.charge_kwh[i] - solution.discharge_kwh[i])
                    / config.dt_hours
                ),
                float(price[i] * 4000.0),
            ]
        )

    detail.freeze_panes = "C5"
    detail.auto_filter.ref = "A4:T148"
    detail.sheet_properties.tabColor = "5B9BD5"
    for cell in detail[4]:
        cell.fill = header_fill
        cell.font = Font(name="Microsoft YaHei", size=9, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    detail.row_dimensions[4].height = 34
    numeric_columns = range(3, 18)
    for row in detail.iter_rows(min_row=5, max_row=148):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=9)
            cell.alignment = Alignment(
                horizontal="center" if cell.column in (1, 2, 18) else "right",
                vertical="center",
            )
            cell.border = Border(bottom=thin_gray)
        for column in numeric_columns:
            row[column - 1].number_format = "#,##0.0000"
    widths = [7, 17, 13, 13, 15, 16, 18, 18, 15, 15, 15, 15, 18, 16, 16, 17, 17, 11, 18, 22]
    for column, width_value in enumerate(widths, start=1):
        detail.column_dimensions[get_column_letter(column)].width = width_value

    charts["A2"] = "问题一：动态规划调度图表"
    charts["A2"].font = Font(name="Microsoft YaHei", size=15, bold=True, color="1F2937")
    charts.sheet_properties.tabColor = "A5A5A5"
    categories = Reference(detail, min_col=2, min_row=5, max_row=148)

    chart1 = LineChart()
    chart1.title = "负荷、光伏与计划购电功率"
    chart1.y_axis.title = "功率/kW"
    chart1.x_axis.title = "时间"
    chart1.height = 9
    chart1.width = 21
    chart1.add_data(Reference(detail, min_col=3, max_col=4, min_row=4, max_row=148), titles_from_data=True)
    chart1.add_data(Reference(detail, min_col=8, min_row=4, max_row=148), titles_from_data=True)
    chart1.set_categories(categories)
    chart1.legend.position = "t"
    charts.add_chart(chart1, "A4")

    chart2 = LineChart()
    chart2.title = "储能充放电功率"
    chart2.y_axis.title = "功率/kW"
    chart2.x_axis.title = "时间"
    chart2.height = 9
    chart2.width = 21
    chart2.add_data(Reference(detail, min_col=11, max_col=12, min_row=4, max_row=148), titles_from_data=True)
    chart2.set_categories(categories)
    chart2.legend.position = "t"
    charts.add_chart(chart2, "L4")

    chart3 = LineChart()
    chart3.title = "储能SOC轨迹"
    chart3.y_axis.title = "储能量/kWh"
    chart3.x_axis.title = "时间"
    chart3.height = 9
    chart3.width = 21
    chart3.add_data(Reference(detail, min_col=14, min_row=4, max_row=148), titles_from_data=True)
    chart3.set_categories(categories)
    chart3.legend = None
    charts.add_chart(chart3, "A22")

    behavior = LineChart()
    behavior.title = "电价与储能行为（电价按4000倍缩放）"
    behavior.y_axis.title = "储能功率/kW；缩放电价"
    behavior.height = 9
    behavior.width = 21
    behavior.add_data(Reference(detail, min_col=19, max_col=20, min_row=4, max_row=148), titles_from_data=True)
    behavior.set_categories(categories)
    charts.add_chart(behavior, "L22")

    for sheet in workbook.worksheets:
        used = sheet.calculate_dimension()
        for row in sheet[used]:
            for cell in row:
                if cell.value is not None and cell.font.name is None:
                    cell.font = Font(name="Microsoft YaHei", size=10)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output


def write_template_result1(
    root: str | Path,
    output_path: str | Path,
    solution: DispatchSolution,
) -> Path:
    """把DP解适配到题目给定的result1.xlsx模板，不调用任何LP求解器。"""
    from microgrid_optimization import DispatchResult, write_problem1_workbook

    adapted = DispatchResult(
        success=True,
        message="dynamic programming solution",
        grid=solution.grid_kwh,
        charge=solution.charge_kwh,
        discharge=solution.discharge_kwh,
        pv_used=np.zeros_like(solution.grid_kwh),
        soc=solution.soc_kwh,
        purchase_cost=float(np.sum(solution.cost_yuan)),
    )
    return write_problem1_workbook(root, output_path, adapted)


def _print_run_report(
    data: dict[str, object],
    solution: DispatchSolution,
    checks: dict[str, bool | float | int],
    config: DispatchConfig,
) -> None:
    print("\n=== 问题一：DP关键结果 ===")
    print(f"全天总购电量: {np.sum(solution.grid_kwh):.6f} kWh")
    print(f"全天总购电费用: {np.sum(solution.cost_yuan):.6f} 元")
    print(f"总充电量: {np.sum(solution.charge_kwh):.6f} kWh（母线侧）")
    print(f"总放电量: {np.sum(solution.discharge_kwh):.6f} kWh（供负荷侧）")
    print(f"最低/最高SOC: {np.min(solution.soc_kwh):.6f} / {np.max(solution.soc_kwh):.6f} kWh")
    print(f"0:00/24:00 SOC: {solution.soc_kwh[0]:.6f} / {solution.soc_kwh[-1]:.6f} kWh")
    print(f"弃光电量: {np.sum(solution.curtailment_kwh):.6f} kWh")
    print(f"最终DP状态分辨率: {solution.resolution_kwh:.6f} kWh")
    print("\n=== 约束检查 ===")
    for key, value in checks.items():
        print(f"{key}: {value}")
    print("\n=== 144个时段最优决策 ===")
    intervals = list(data["intervals"])
    for i, interval in enumerate(intervals):
        if solution.charge_kwh[i] > config.tolerance:
            action = f"充电 {solution.charge_kwh[i]:.4f} kWh"
        elif solution.discharge_kwh[i] > config.tolerance:
            action = f"放电 {solution.discharge_kwh[i]:.4f} kWh"
        else:
            action = "不动作"
        print(
            f"{i + 1:03d} {interval} | {action} | "
            f"购电 {solution.grid_kwh[i]:.4f} kWh | "
            f"SOC {solution.soc_kwh[i + 1]:.4f} kWh | "
            f"费用 {solution.cost_yuan[i]:.4f} 元"
        )


def main() -> None:
    import argparse
    import time as time_module

    parser = argparse.ArgumentParser(description="使用动态规划独立求解微网问题一")
    parser.add_argument("--input", default=str(Path("附件") / "附件1.xlsx"))
    parser.add_argument("--output", default="result_q1.xlsx")
    parser.add_argument("--template-output", default="result1.xlsx")
    parser.add_argument("--plot-dir", default="q1_dp_figures")
    args = parser.parse_args()

    config = DispatchConfig()
    data = read_problem1_data(args.input)
    started = time_module.perf_counter()
    solution = run_adaptive_dp(
        np.asarray(data["load_kw"]),
        np.asarray(data["pv_kw"]),
        np.asarray(data["price"]),
        config,
    )
    elapsed = time_module.perf_counter() - started
    checks = validate_dispatch(
        np.asarray(data["load_kw"]),
        np.asarray(data["pv_kw"]),
        np.asarray(data["price"]),
        solution.soc_kwh,
        solution.grid_kwh,
        solution.charge_kwh,
        solution.discharge_kwh,
        solution.curtailment_kwh,
        config,
    )
    failed = [key for key, value in checks.items() if isinstance(value, bool) and not value]
    if failed:
        raise RuntimeError(f"约束检查失败：{failed}")
    workbook_path = write_result_workbook(args.output, data, solution, checks, config)
    template_path = write_template_result1(
        Path(__file__).resolve().parent, args.template_output, solution
    )
    plot_paths = write_plots(args.plot_dir, data, solution, config)
    _print_run_report(data, solution, checks, config)
    print(f"\nDP运行耗时: {elapsed:.3f} s")
    print(f"结果文件: {workbook_path.resolve()}")
    print(f"题目模板结果: {template_path.resolve()}")
    for plot_path in plot_paths:
        print(f"图像文件: {plot_path.resolve()}")


if __name__ == "__main__":
    main()

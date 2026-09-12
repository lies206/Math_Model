"""2026 数学建模 C 题问题三：多阶段预测、MPC 与非对称调整成本。

本模块的数据入口严格白名单化：只读取附件1.xlsx、附件2.xlsx、附件3.xlsx。
附件2的144点按10分钟区间终点记录；内部统一表示为从0:00开始的144个区间电量。
附件3每条“预测k小时”解释为发布时间后第k个整点的功率预测，并以发布时间
已经观测到的实际光伏为左端锚点，线性插值到10分钟区间终点。
"""

from __future__ import annotations

from dataclasses import asdict, dataclass
import json
from pathlib import Path
from typing import Sequence

import numpy as np

from microgrid_optimization import (
    DispatchConfig,
    DispatchResult,
    RollingYearResult,
    audit_dispatch,
    kw_to_interval_kwh,
    solve_dispatch,
    write_rolling_workbook,
)
from problem2_optimization import price_for_day


OUTPUT_START_DAY = 31  # 2025-02-01，1月仅作因果预热
RELEASE_HOURS = (0, 6, 12, 18)
LOAD_WINDOWS = (7, 14, 21, 28)
LOAD_DECAYS = (0.65, 0.80, 0.90, 1.00)
VALIDATION_DAYS = 28
TOLERANCE = 1e-5


@dataclass(frozen=True)
class Problem3Data:
    static_price: np.ndarray
    dates: np.ndarray
    load_kw: np.ndarray
    actual_pv_kw: np.ndarray
    pv_forecast_hourly_kw: np.ndarray
    source_files: tuple[str, str, str]


@dataclass(frozen=True)
class LoadForecastResult:
    forecast_kw: np.ndarray
    selected_config: np.ndarray
    validation_rmse_kw: np.ndarray
    max_history_day_used: np.ndarray

    @property
    def leakage_detected(self) -> bool:
        return bool(np.any(self.max_history_day_used >= np.arange(len(self.max_history_day_used))))


@dataclass(frozen=True)
class AdjustmentCost:
    downward_kwh: float
    upward_kwh: float
    downward_cost_yuan: float
    upward_cost_yuan: float

    @property
    def total_cost_yuan(self) -> float:
        return self.downward_cost_yuan + self.upward_cost_yuan


@dataclass(frozen=True)
class RealtimeSegment:
    grid_used: np.ndarray
    charge: np.ndarray
    discharge: np.ndarray
    pv_used: np.ndarray
    pv_curtailed: np.ndarray
    unused_plan: np.ndarray
    emergency: np.ndarray
    soc: np.ndarray


@dataclass(frozen=True)
class DayStrategyResult:
    plan_grid: np.ndarray
    adjusted_grid: np.ndarray
    grid_used: np.ndarray
    charge: np.ndarray
    discharge: np.ndarray
    pv_used: np.ndarray
    pv_curtailed: np.ndarray
    unused_plan: np.ndarray
    emergency: np.ndarray
    soc: np.ndarray
    plan_cost_yuan: float
    adjustment_cost_yuan: float
    downward_cost_yuan: float
    upward_cost_yuan: float
    emergency_cost_yuan: float
    downward_kwh: float
    upward_kwh: float
    adjustment_interval_count: int
    adjustment_event_count: int
    update_records: list[UpdateRecord]


@dataclass(frozen=True)
class UpdateRecord:
    day_index: int
    release_hour: int
    start_slot: int
    inherited_actual_soc_kwh: float
    downward_kwh: float
    upward_kwh: float
    downward_cost_yuan: float
    upward_cost_yuan: float
    expected_keep_emergency_cost_yuan: float
    expected_revised_emergency_cost_yuan: float
    expected_gross_saving_yuan: float
    applied: bool


@dataclass
class StrategyResult:
    name: str
    update_hours: tuple[int, ...]
    event_triggered: bool
    plan_grid: np.ndarray
    adjusted_grid: np.ndarray
    grid_used: np.ndarray
    charge: np.ndarray
    discharge: np.ndarray
    pv_used: np.ndarray
    pv_curtailed: np.ndarray
    unused_plan: np.ndarray
    emergency: np.ndarray
    soc: np.ndarray
    daily_plan_cost: np.ndarray
    daily_adjustment_cost: np.ndarray
    daily_downward_cost: np.ndarray
    daily_upward_cost: np.ndarray
    daily_emergency_cost: np.ndarray
    daily_downward_kwh: np.ndarray
    daily_upward_kwh: np.ndarray
    daily_adjustment_interval_count: np.ndarray
    daily_adjustment_event_count: np.ndarray
    update_records: list[UpdateRecord]

    @property
    def daily_total_cost(self) -> np.ndarray:
        return self.daily_plan_cost + self.daily_adjustment_cost + self.daily_emergency_cost


@dataclass(frozen=True)
class Problem3Study:
    selected_result: StrategyResult
    strategies: dict[str, StrategyResult]
    strategy_metrics: dict[str, dict[str, float | int | bool]]
    pv_accuracy: dict[str, dict[str, float | int]]
    marginal_information_value: dict[str, dict[str, float | bool]]
    load_forecast: LoadForecastResult
    audit: dict[str, object]


def allowed_input_files(root: str | Path) -> tuple[Path, Path, Path]:
    """返回唯一允许的问题三数据源；该函数不会探测或打开附件4。"""
    folder = Path(root).resolve() / "附件"
    paths = tuple(folder / name for name in ("附件1.xlsx", "附件2.xlsx", "附件3.xlsx"))
    missing = [path.name for path in paths if not path.is_file()]
    if missing:
        raise FileNotFoundError(f"问题三缺少输入文件: {missing}")
    return paths  # type: ignore[return-value]


def _describe_workbook(path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    print(f"{path.name}: sheets={workbook.sheetnames}")
    for sheet in workbook.worksheets:
        header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        print(f"  {sheet.title}: shape=({sheet.max_row}, {sheet.max_column}), columns={header[:6]}")
    workbook.close()


def _find_sheet(workbook, keyword: str):
    matches = [sheet for sheet in workbook.worksheets if keyword in sheet.title]
    if len(matches) != 1:
        raise ValueError(f"无法唯一识别包含‘{keyword}’的工作表")
    return matches[0]


def _parse_date(value: object) -> np.datetime64:
    from datetime import datetime

    if hasattr(value, "date"):
        value = value.date()
    text = str(value).strip()
    for pattern in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return np.datetime64(datetime.strptime(text, pattern).date(), "D")
        except ValueError:
            continue
    raise ValueError(f"无法解析日期: {value}")


def load_problem3_data(root: str | Path, print_structure: bool = True) -> Problem3Data:
    """严格读取附件1电价、附件2实际负荷/光伏、附件3分时光伏预报。"""
    from openpyxl import load_workbook

    price_path, actual_path, forecast_path = allowed_input_files(root)
    if print_structure:
        for path in (price_path, actual_path, forecast_path):
            _describe_workbook(path)

    workbook1 = load_workbook(price_path, read_only=True, data_only=True)
    sheet1 = workbook1.worksheets[0]
    headers = [cell.value for cell in next(sheet1.iter_rows(min_row=1, max_row=1))]
    price_columns = [index + 1 for index, value in enumerate(headers) if "电价" in str(value)]
    if len(price_columns) != 1:
        raise ValueError("附件1无法唯一识别电价列")
    price = np.asarray(
        [row[0] for row in sheet1.iter_rows(min_row=2, min_col=price_columns[0], max_col=price_columns[0], values_only=True)],
        dtype=float,
    )
    workbook1.close()

    workbook2 = load_workbook(actual_path, read_only=True, data_only=True)
    load_sheet = _find_sheet(workbook2, "负载")
    pv_sheet = _find_sheet(workbook2, "实际")
    load_rows = list(load_sheet.iter_rows(min_row=2, values_only=True))
    pv_rows = list(pv_sheet.iter_rows(min_row=2, values_only=True))
    workbook2.close()
    dates = np.asarray([_parse_date(row[0]) for row in load_rows], dtype="datetime64[D]")
    pv_dates = np.asarray([_parse_date(row[0]) for row in pv_rows], dtype="datetime64[D]")
    load_kw = np.asarray([row[1:] for row in load_rows], dtype=float)
    actual_pv_kw = np.asarray([row[1:] for row in pv_rows], dtype=float)

    workbook3 = load_workbook(forecast_path, read_only=True, data_only=True)
    forecast_sheet = workbook3.worksheets[0]
    forecast_rows = list(forecast_sheet.iter_rows(min_row=2, values_only=True))
    workbook3.close()
    if len(forecast_rows) != 365 * 4:
        raise ValueError(f"附件3应有1460条预报，实际为{len(forecast_rows)}")
    forecast_dates: list[np.datetime64] = []
    issue_hours: list[int] = []
    forecast_values: list[tuple[float, ...]] = []
    current_date: np.datetime64 | None = None
    for row in forecast_rows:
        if row[0] not in (None, ""):
            current_date = _parse_date(row[0])
        if current_date is None:
            raise ValueError("附件3首条预报缺少日期")
        issue_text = str(row[1]).strip()
        try:
            issue_hour = int(issue_text.split(":")[0])
        except (TypeError, ValueError, IndexError) as exc:
            raise ValueError(f"附件3非法预报时刻: {row[1]}") from exc
        forecast_dates.append(current_date)
        issue_hours.append(issue_hour)
        forecast_values.append(tuple(float(value) for value in row[2:26]))
    forecasts = np.asarray(forecast_values, dtype=float).reshape(365, 4, 24)
    forecast_date_matrix = np.asarray(forecast_dates, dtype="datetime64[D]").reshape(365, 4)
    issue_matrix = np.asarray(issue_hours, dtype=int).reshape(365, 4)

    expected_dates = np.arange(np.datetime64("2025-01-01"), np.datetime64("2026-01-01"))
    if price.shape != (144,):
        raise ValueError(f"附件1电价应为144点，实际为{price.shape}")
    if load_kw.shape != (365, 144) or actual_pv_kw.shape != (365, 144):
        raise ValueError(f"附件2应为365×144，实际负荷{load_kw.shape}、光伏{actual_pv_kw.shape}")
    if not np.array_equal(dates, expected_dates) or not np.array_equal(pv_dates, expected_dates):
        raise ValueError("附件2日期不连续或负荷/光伏日期不一致")
    if not np.array_equal(forecast_date_matrix, np.repeat(expected_dates[:, None], 4, axis=1)):
        raise ValueError("附件3日期与附件2不一致")
    if not np.array_equal(issue_matrix, np.repeat(np.asarray(RELEASE_HOURS)[None, :], 365, axis=0)):
        raise ValueError("附件3每天必须依次包含0:00、6:00、12:00、18:00预报")
    for name, array in (("电价", price), ("负荷", load_kw), ("实际光伏", actual_pv_kw), ("光伏预报", forecasts)):
        if not np.all(np.isfinite(array)) or np.any(array < 0.0):
            raise ValueError(f"{name}包含NaN/Inf或负值")
    return Problem3Data(
        price,
        dates,
        load_kw,
        actual_pv_kw,
        forecasts,
        tuple(path.name for path in (price_path, actual_path, forecast_path)),
    )


def _same_weekday_forecast(actual_kw: np.ndarray, day: int, window_days: int, decay: float) -> np.ndarray:
    lags = [lag for lag in range(7, window_days + 1, 7) if day - lag >= 0]
    if not lags:
        if day == 0:
            return np.zeros(actual_kw.shape[1], dtype=float)
        return actual_kw[day - 1].copy()
    weights = decay ** np.arange(len(lags), dtype=float)
    weights /= weights.sum()
    return np.sum(np.asarray([actual_kw[day - lag] for lag in lags]) * weights[:, None], axis=0)


def build_causal_load_forecasts(
    actual_load_kw: np.ndarray,
    windows: Sequence[int] = LOAD_WINDOWS,
    decays: Sequence[float] = LOAD_DECAYS,
    validation_days: int = VALIDATION_DAYS,
) -> LoadForecastResult:
    """同星期同时刻预测；参数只能由当前日前已经结束的滚动验证集选择。"""
    actual = np.asarray(actual_load_kw, dtype=float)
    if actual.ndim != 2 or actual.shape[1] != 144:
        raise ValueError("负荷数据必须为days×144")
    if not np.all(np.isfinite(actual)) or np.any(actual < 0.0):
        raise ValueError("负荷数据包含NaN/Inf或负值")
    configs = [(int(window), float(decay)) for window in windows for decay in decays]
    if not configs:
        raise ValueError("候选窗口和衰减参数不能为空")
    if any(window not in LOAD_WINDOWS for window, _ in configs):
        raise ValueError("窗口必须从7/14/21/28日中选择")
    if any(not 0.0 < decay <= 1.0 for _, decay in configs):
        raise ValueError("衰减参数必须位于(0,1]")

    days = len(actual)
    candidates = np.zeros((len(configs), days, 144), dtype=float)
    squared_error = np.full((len(configs), days), np.nan)
    for day in range(days):
        for index, (window, decay) in enumerate(configs):
            forecast = _same_weekday_forecast(actual, day, window, decay)
            candidates[index, day] = forecast
            squared_error[index, day] = float(np.mean((forecast - actual[day]) ** 2))

    forecast = np.zeros_like(actual)
    selected = np.empty(days, dtype=object)
    validation_rmse = np.full(days, np.nan)
    max_history = np.full(days, -1, dtype=int)
    default_index = configs.index((28, 0.80)) if (28, 0.80) in configs else 0
    for day in range(days):
        start = max(0, day - validation_days)
        # 第day日的选择仅使用start..day-1日已经发生的预测误差。
        if day - start >= 7:
            score = np.sqrt(np.mean(squared_error[:, start:day], axis=1))
            chosen = int(np.argmin(score))
            validation_rmse[day] = float(score[chosen])
        else:
            chosen = default_index
        window, decay = configs[chosen]
        forecast[day] = candidates[chosen, day]
        selected[day] = f"w{window}_decay{decay:.2f}"
        max_history[day] = day - 7 if day >= 7 else day - 1
    return LoadForecastResult(forecast, selected, validation_rmse, max_history)


def online_load_bias_correction(
    base_forecast_kw: np.ndarray | Sequence[float],
    actual_load_kw: np.ndarray | Sequence[float],
    start_slot: int,
    lookback_slots: int = 18,
    decay: float = 0.88,
) -> np.ndarray:
    """用发布时点之前至多3小时残差修正剩余负荷；不读取start_slot及以后实际值。"""
    base = np.asarray(base_forecast_kw, dtype=float).reshape(-1)
    actual = np.asarray(actual_load_kw, dtype=float).reshape(-1)
    if base.shape != (144,) or actual.shape != (144,):
        raise ValueError("单日负荷预测和实际值均须为144点")
    if not 0 <= start_slot <= 144:
        raise ValueError("start_slot越界")
    corrected = base.copy()
    if start_slot == 0:
        return np.maximum(corrected, 0.0)
    begin = max(0, start_slot - lookback_slots)
    residual = actual[begin:start_slot] - base[begin:start_slot]
    weights = decay ** np.arange(len(residual) - 1, -1, -1, dtype=float)
    bias = float(np.dot(weights, residual) / weights.sum())
    corrected[start_slot:] = np.maximum(base[start_slot:] + bias, 0.0)
    return corrected


def interpolate_pv_release(
    hourly_forecast_kw: np.ndarray | Sequence[float],
    release_hour: int,
    anchor_kw: float,
) -> np.ndarray:
    """将发布时刻后的整点功率线性插值到当天剩余10分钟区间终点。"""
    hourly = np.asarray(hourly_forecast_kw, dtype=float).reshape(-1)
    if hourly.shape != (24,):
        raise ValueError("每个发布时间必须包含未来24个整点预测")
    if release_hour not in RELEASE_HOURS or anchor_kw < 0.0:
        raise ValueError("非法发布时间或光伏锚点")
    remaining_hours = 24 - release_hour
    anchors = np.r_[float(anchor_kw), hourly[:remaining_hours]]
    source_hours = np.arange(remaining_hours + 1, dtype=float)
    target_hours = np.arange(1, remaining_hours * 6 + 1, dtype=float) / 6.0
    return np.maximum(np.interp(target_hours, source_hours, anchors), 0.0)


def asymmetric_adjustment_cost(
    reference_grid_kwh: np.ndarray | Sequence[float],
    revised_grid_kwh: np.ndarray | Sequence[float],
    price: np.ndarray | Sequence[float],
) -> AdjustmentCost:
    """分别核算下调违约费0.5c与上调超购费1.5c，绝不合并为绝对值。"""
    reference = np.asarray(reference_grid_kwh, dtype=float)
    revised = np.asarray(revised_grid_kwh, dtype=float)
    prices = np.asarray(price, dtype=float)
    if not (reference.shape == revised.shape == prices.shape):
        raise ValueError("调整基准、调整结果和电价必须同形")
    downward = np.maximum(reference - revised, 0.0)
    upward = np.maximum(revised - reference, 0.0)
    return AdjustmentCost(
        float(downward.sum()),
        float(upward.sum()),
        float(np.dot(0.5 * prices, downward)),
        float(np.dot(1.5 * prices, upward)),
    )


def _solve_milp_schedule(
    load_kwh: np.ndarray,
    pv_kwh: np.ndarray,
    price: np.ndarray,
    initial_soc: float,
    terminal_soc: float,
    reference_grid_kwh: np.ndarray | None,
    config: DispatchConfig,
) -> DispatchResult:
    from scipy.optimize import Bounds, LinearConstraint, milp
    from scipy.sparse import coo_matrix

    load = np.asarray(load_kwh, dtype=float).reshape(-1)
    pv = np.asarray(pv_kwh, dtype=float).reshape(-1)
    prices = np.asarray(price, dtype=float).reshape(-1)
    periods = len(load)
    if periods == 0 or not (load.shape == pv.shape == prices.shape):
        raise ValueError("MPC负荷、光伏和电价必须为非空同形数组")
    if np.any(load < 0.0) or np.any(pv < 0.0) or np.any(prices < 0.0):
        raise ValueError("MPC输入不能为负")
    if not config.soc_min_kwh <= initial_soc <= config.soc_max_kwh:
        raise ValueError("MPC当前实际SOC越界")
    if not config.soc_min_kwh <= terminal_soc <= config.soc_max_kwh:
        raise ValueError("MPC日末SOC目标越界")
    reference = None if reference_grid_kwh is None else np.asarray(reference_grid_kwh, dtype=float).reshape(-1)
    if reference is not None and reference.shape != load.shape:
        raise ValueError("调整基准与剩余时域长度不一致")

    grid_offset = 0
    charge_offset = periods
    discharge_offset = 2 * periods
    pv_offset = 3 * periods
    soc_offset = 4 * periods
    soc_end = 5 * periods + 1
    if reference is None:
        down_offset = up_offset = mode_offset = soc_end
        variable_count = soc_end + periods
    else:
        down_offset = soc_end
        up_offset = down_offset + periods
        mode_offset = up_offset + periods
        variable_count = mode_offset + periods

    objective = np.zeros(variable_count)
    if reference is None:
        objective[grid_offset:charge_offset] = prices
    else:
        objective[down_offset:up_offset] = 0.5 * prices
        objective[up_offset:mode_offset] = 1.5 * prices
    objective[charge_offset:discharge_offset] = config.throughput_penalty
    objective[discharge_offset:pv_offset] = config.throughput_penalty

    rows: list[int] = []
    columns: list[int] = []
    values: list[float] = []
    lower_rows: list[float] = []
    upper_rows: list[float] = []
    row_index = 0
    for t in range(periods):
        rows.extend((row_index,) * 4)
        columns.extend((grid_offset + t, pv_offset + t, discharge_offset + t, charge_offset + t))
        values.extend((1.0, 1.0, 1.0, -1.0))
        lower_rows.append(load[t])
        upper_rows.append(load[t])
        row_index += 1

        rows.extend((row_index,) * 4)
        columns.extend((soc_offset + t, soc_offset + t + 1, charge_offset + t, discharge_offset + t))
        values.extend((-1.0, 1.0, -config.efficiency, 1.0 / config.efficiency))
        lower_rows.append(0.0)
        upper_rows.append(0.0)
        row_index += 1

        if reference is not None:
            rows.extend((row_index,) * 3)
            columns.extend((grid_offset + t, down_offset + t, up_offset + t))
            values.extend((1.0, 1.0, -1.0))
            lower_rows.append(reference[t])
            upper_rows.append(reference[t])
            row_index += 1

        rows.extend((row_index, row_index))
        columns.extend((charge_offset + t, mode_offset + t))
        values.extend((1.0, -config.interval_limit_kwh))
        lower_rows.append(-np.inf)
        upper_rows.append(0.0)
        row_index += 1

        rows.extend((row_index, row_index))
        columns.extend((discharge_offset + t, mode_offset + t))
        values.extend((1.0, config.interval_limit_kwh))
        lower_rows.append(-np.inf)
        upper_rows.append(config.interval_limit_kwh)
        row_index += 1

    matrix = coo_matrix((values, (rows, columns)), shape=(row_index, variable_count)).tocsr()
    lower = np.zeros(variable_count)
    upper = np.full(variable_count, np.inf)
    upper[charge_offset:discharge_offset] = config.interval_limit_kwh
    upper[discharge_offset:pv_offset] = config.interval_limit_kwh
    upper[pv_offset:soc_offset] = pv
    lower[soc_offset:soc_end] = config.soc_min_kwh
    upper[soc_offset:soc_end] = config.soc_max_kwh
    lower[soc_offset] = upper[soc_offset] = float(initial_soc)
    lower[soc_end - 1] = upper[soc_end - 1] = float(terminal_soc)
    upper[mode_offset:mode_offset + periods] = 1.0
    integrality = np.zeros(variable_count, dtype=int)
    integrality[mode_offset:mode_offset + periods] = 1

    result = milp(
        objective,
        integrality=integrality,
        bounds=Bounds(lower, upper),
        constraints=LinearConstraint(matrix, np.asarray(lower_rows), np.asarray(upper_rows)),
        options={"presolve": True},
    )
    if not result.success:
        empty = np.full(periods, np.nan)
        return DispatchResult(False, result.message, empty, empty.copy(), empty.copy(), empty.copy(), np.full(periods + 1, np.nan), np.nan)
    solution = result.x
    return DispatchResult(
        True,
        result.message,
        solution[grid_offset:charge_offset],
        solution[charge_offset:discharge_offset],
        solution[discharge_offset:pv_offset],
        solution[pv_offset:soc_offset],
        solution[soc_offset:soc_end],
        float(np.dot(prices, solution[grid_offset:charge_offset])),
    )


def solve_mpc_schedule(
    forecast_load_kwh: np.ndarray | Sequence[float],
    forecast_pv_kwh: np.ndarray | Sequence[float],
    price: np.ndarray | Sequence[float],
    initial_soc: float,
    terminal_soc: float,
    reference_grid_kwh: np.ndarray | Sequence[float] | None = None,
    config: DispatchConfig = DispatchConfig(),
) -> DispatchResult:
    """求解剩余时域；有reference时目标仅为方向分离的附加调整费。"""
    load = np.asarray(forecast_load_kwh, dtype=float).reshape(-1)
    pv = np.asarray(forecast_pv_kwh, dtype=float).reshape(-1)
    prices = np.asarray(price, dtype=float).reshape(-1)
    reference = None if reference_grid_kwh is None else np.asarray(reference_grid_kwh, dtype=float).reshape(-1)
    # 日前问题是普通电价下的凸LP；正的吞吐惩罚排除退化的同时充放电，速度远快于全年MILP。
    if reference is None:
        result = solve_dispatch(load, pv, prices, initial_soc, terminal_soc, config)
        if result.success and np.max(np.minimum(result.charge, result.discharge)) <= TOLERANCE:
            return result
    return _solve_milp_schedule(load, pv, prices, initial_soc, terminal_soc, reference, config)


def simulate_realtime_segment(
    contracted_grid_kwh: np.ndarray | Sequence[float],
    actual_load_kwh: np.ndarray | Sequence[float],
    actual_pv_kwh: np.ndarray | Sequence[float],
    initial_soc: float,
    config: DispatchConfig = DispatchConfig(),
) -> RealtimeSegment:
    """执行冻结区间：先用实际光伏和合同购电，再充/放电，最后才紧急购电。"""
    grid = np.asarray(contracted_grid_kwh, dtype=float).reshape(-1)
    load = np.asarray(actual_load_kwh, dtype=float).reshape(-1)
    pv = np.asarray(actual_pv_kwh, dtype=float).reshape(-1)
    if len(grid) == 0 or not (grid.shape == load.shape == pv.shape):
        raise ValueError("实时执行数组必须为非空同形")
    if np.any(grid < -TOLERANCE) or np.any(load < 0.0) or np.any(pv < 0.0):
        raise ValueError("实时执行输入不能为负")
    if not config.soc_min_kwh <= initial_soc <= config.soc_max_kwh:
        raise ValueError("实时执行当前SOC越界")

    periods = len(grid)
    grid_used = np.zeros(periods)
    charge = np.zeros(periods)
    discharge = np.zeros(periods)
    pv_used = np.zeros(periods)
    curtailed = np.zeros(periods)
    unused = np.zeros(periods)
    emergency = np.zeros(periods)
    soc = np.zeros(periods + 1)
    soc[0] = float(initial_soc)
    for slot in range(periods):
        available = grid[slot] + pv[slot]
        if available >= load[slot]:
            surplus = available - load[slot]
            headroom_input = max((config.soc_max_kwh - soc[slot]) / config.efficiency, 0.0)
            charge[slot] = min(surplus, config.interval_limit_kwh, headroom_input)
            bus_need = load[slot] + charge[slot]
            pv_used[slot] = min(pv[slot], bus_need)
            grid_used[slot] = min(grid[slot], max(bus_need - pv_used[slot], 0.0))
            curtailed[slot] = pv[slot] - pv_used[slot]
            unused[slot] = grid[slot] - grid_used[slot]
        else:
            grid_used[slot] = grid[slot]
            pv_used[slot] = pv[slot]
            deficit = load[slot] - available
            available_output = max((soc[slot] - config.soc_min_kwh) * config.efficiency, 0.0)
            discharge[slot] = min(deficit, config.interval_limit_kwh, available_output)
            emergency[slot] = deficit - discharge[slot]
        soc[slot + 1] = soc[slot] + config.efficiency * charge[slot] - discharge[slot] / config.efficiency

    balance = grid_used + pv_used + discharge + emergency - load - charge
    if np.max(np.abs(balance)) > TOLERANCE:
        raise RuntimeError("实时执行电量平衡失败")
    if soc.min() < config.soc_min_kwh - TOLERANCE or soc.max() > config.soc_max_kwh + TOLERANCE:
        raise RuntimeError("实时执行SOC越界")
    if np.max(np.minimum(charge, discharge)) > TOLERANCE:
        raise RuntimeError("实时执行出现同时充放电")
    return RealtimeSegment(grid_used, charge, discharge, pv_used, curtailed, unused, emergency, soc)


def _pv_anchor_kw(day_index: int, release_hour: int, actual_pv_kw: np.ndarray) -> float:
    """发布时间的可用锚点：0点用前一日24点，日内用刚结束区间的实测值。"""
    if release_hour == 0:
        return 0.0 if day_index == 0 else float(actual_pv_kw[day_index - 1, -1])
    return float(actual_pv_kw[day_index, release_hour * 6 - 1])


def run_day_strategy(
    day_index: int,
    actual_load_kw: np.ndarray | Sequence[float],
    actual_pv_kw: np.ndarray | Sequence[float],
    base_load_forecast_kw: np.ndarray | Sequence[float],
    pv_releases_hourly_kw: np.ndarray,
    price: np.ndarray | Sequence[float],
    initial_soc: float,
    update_hours: Sequence[int],
    event_triggered: bool,
    previous_day_pv_end_kw: float = 0.0,
    config: DispatchConfig = DispatchConfig(),
) -> DayStrategyResult:
    """运行单日策略；每次更新前先执行并冻结历史，再用真实SOC求解剩余时域。"""
    actual_load_power = np.asarray(actual_load_kw, dtype=float).reshape(-1)
    actual_pv_power = np.asarray(actual_pv_kw, dtype=float).reshape(-1)
    base_load_power = np.asarray(base_load_forecast_kw, dtype=float).reshape(-1)
    releases = np.asarray(pv_releases_hourly_kw, dtype=float)
    prices = np.asarray(price, dtype=float).reshape(-1)
    if not (
        actual_load_power.shape == actual_pv_power.shape == base_load_power.shape == prices.shape == (144,)
    ):
        raise ValueError("单日负荷、光伏、预测和电价必须均为144点")
    if releases.shape != (4, 24):
        raise ValueError("单日附件3预报必须为4×24")
    updates = tuple(sorted(set(int(hour) for hour in update_hours)))
    if any(hour not in (6, 12, 18) for hour in updates):
        raise ValueError("日内更新只能发生在6:00、12:00、18:00")

    actual_load_energy = kw_to_interval_kwh(actual_load_power, config.interval_hours)
    actual_pv_energy = kw_to_interval_kwh(actual_pv_power, config.interval_hours)
    load0_energy = kw_to_interval_kwh(np.maximum(base_load_power, 0.0), config.interval_hours)
    pv0_energy = kw_to_interval_kwh(
        interpolate_pv_release(releases[0], 0, max(float(previous_day_pv_end_kw), 0.0)),
        config.interval_hours,
    )
    daily_terminal_soc = float(initial_soc)
    plan = solve_mpc_schedule(
        load0_energy,
        pv0_energy,
        prices,
        initial_soc,
        daily_terminal_soc,
        reference_grid_kwh=None,
        config=config,
    )
    if not plan.success:
        raise RuntimeError(f"第{day_index + 1}日0:00计划求解失败: {plan.message}")

    active_grid = plan.grid.copy()
    adjusted_grid = plan.grid.copy()
    grid_used = np.zeros(144)
    charge = np.zeros(144)
    discharge = np.zeros(144)
    pv_used = np.zeros(144)
    pv_curtailed = np.zeros(144)
    unused_plan = np.zeros(144)
    emergency = np.zeros(144)
    soc = np.zeros(145)
    soc[0] = float(initial_soc)
    adjustment_cost = downward_cost = upward_cost = 0.0
    downward_kwh = upward_kwh = 0.0
    adjustment_interval_count = adjustment_event_count = 0
    update_records: list[UpdateRecord] = []
    current_slot = 0

    for release_hour in (*updates, 24):
        start_slot = release_hour * 6
        if start_slot > current_slot:
            executed = simulate_realtime_segment(
                active_grid[current_slot:start_slot],
                actual_load_energy[current_slot:start_slot],
                actual_pv_energy[current_slot:start_slot],
                float(soc[current_slot]),
                config,
            )
            grid_used[current_slot:start_slot] = executed.grid_used
            charge[current_slot:start_slot] = executed.charge
            discharge[current_slot:start_slot] = executed.discharge
            pv_used[current_slot:start_slot] = executed.pv_used
            pv_curtailed[current_slot:start_slot] = executed.pv_curtailed
            unused_plan[current_slot:start_slot] = executed.unused_plan
            emergency[current_slot:start_slot] = executed.emergency
            soc[current_slot : start_slot + 1] = executed.soc
            current_slot = start_slot
        if release_hour == 24:
            break

        # 此处actual_load_power[:start_slot]以及actual_pv_power[start_slot-1]均已观测；
        # start_slot及以后实际数据不进入预测或优化。
        corrected_load_power = online_load_bias_correction(
            base_load_power, actual_load_power, start_slot
        )
        remaining_load_energy = kw_to_interval_kwh(
            corrected_load_power[start_slot:], config.interval_hours
        )
        release_index = RELEASE_HOURS.index(release_hour)
        anchor = float(actual_pv_power[start_slot - 1])
        remaining_pv_energy = kw_to_interval_kwh(
            interpolate_pv_release(releases[release_index], release_hour, anchor),
            config.interval_hours,
        )
        reference = active_grid[start_slot:].copy()
        inherited_soc = float(soc[start_slot])
        revised = solve_mpc_schedule(
            remaining_load_energy,
            remaining_pv_energy,
            prices[start_slot:],
            inherited_soc,
            daily_terminal_soc,
            reference_grid_kwh=reference,
            config=config,
        )
        if not revised.success:
            raise RuntimeError(f"第{day_index + 1}日{release_hour}:00滚动优化失败: {revised.message}")
        settlement = asymmetric_adjustment_cost(reference, revised.grid, prices[start_slot:])

        expected_keep = simulate_realtime_segment(
            reference,
            remaining_load_energy,
            remaining_pv_energy,
            inherited_soc,
            config,
        )
        expected_revised = simulate_realtime_segment(
            revised.grid,
            remaining_load_energy,
            remaining_pv_energy,
            inherited_soc,
            config,
        )
        expected_keep_emergency_cost = float(np.dot(5.0 * prices[start_slot:], expected_keep.emergency))
        expected_revised_emergency_cost = float(np.dot(5.0 * prices[start_slot:], expected_revised.emergency))
        expected_gross_saving = expected_keep_emergency_cost - expected_revised_emergency_cost
        apply_update = (not event_triggered) or (
            expected_gross_saving > settlement.total_cost_yuan + 1e-7
        )
        update_records.append(
            UpdateRecord(
                day_index,
                release_hour,
                start_slot,
                inherited_soc,
                settlement.downward_kwh,
                settlement.upward_kwh,
                settlement.downward_cost_yuan,
                settlement.upward_cost_yuan,
                expected_keep_emergency_cost,
                expected_revised_emergency_cost,
                expected_gross_saving,
                apply_update,
            )
        )
        if apply_update:
            delta = np.abs(revised.grid - reference)
            if np.any(delta > TOLERANCE):
                adjustment_event_count += 1
                adjustment_interval_count += int(np.sum(delta > TOLERANCE))
            adjustment_cost += settlement.total_cost_yuan
            downward_cost += settlement.downward_cost_yuan
            upward_cost += settlement.upward_cost_yuan
            downward_kwh += settlement.downward_kwh
            upward_kwh += settlement.upward_kwh
            # 只覆盖尚未执行的后缀，前缀同时用断言证明已冻结。
            frozen_prefix = active_grid[:start_slot].copy()
            active_grid[start_slot:] = revised.grid
            adjusted_grid[start_slot:] = revised.grid
            if not np.array_equal(active_grid[:start_slot], frozen_prefix):
                raise RuntimeError("滚动优化改写了已执行时段")

    emergency_cost = float(np.dot(5.0 * prices, emergency))
    return DayStrategyResult(
        plan.grid,
        adjusted_grid,
        grid_used,
        charge,
        discharge,
        pv_used,
        pv_curtailed,
        unused_plan,
        emergency,
        soc,
        float(np.dot(prices, plan.grid)),
        adjustment_cost,
        downward_cost,
        upward_cost,
        emergency_cost,
        downward_kwh,
        upward_kwh,
        adjustment_interval_count,
        adjustment_event_count,
        update_records,
    )


def run_strategy(
    name: str,
    data: Problem3Data,
    load_forecast: LoadForecastResult,
    update_hours: Sequence[int],
    event_triggered: bool = False,
    config: DispatchConfig = DispatchConfig(),
    progress: bool = True,
) -> StrategyResult:
    days, periods = data.load_kw.shape
    shape = (days, periods)
    plan_grid = np.zeros(shape)
    adjusted_grid = np.zeros(shape)
    grid_used = np.zeros(shape)
    charge = np.zeros(shape)
    discharge = np.zeros(shape)
    pv_used = np.zeros(shape)
    pv_curtailed = np.zeros(shape)
    unused_plan = np.zeros(shape)
    emergency = np.zeros(shape)
    soc = np.zeros((days, periods + 1))
    daily_plan_cost = np.zeros(days)
    daily_adjustment_cost = np.zeros(days)
    daily_downward_cost = np.zeros(days)
    daily_upward_cost = np.zeros(days)
    daily_emergency_cost = np.zeros(days)
    daily_downward_kwh = np.zeros(days)
    daily_upward_kwh = np.zeros(days)
    daily_adjustment_interval_count = np.zeros(days, dtype=int)
    daily_adjustment_event_count = np.zeros(days, dtype=int)
    records: list[UpdateRecord] = []
    current_soc = 6000.0
    for day in range(days):
        previous_pv_end = 0.0 if day == 0 else float(data.actual_pv_kw[day - 1, -1])
        day_price = price_for_day(data.static_price, day, days)
        result = run_day_strategy(
            day,
            data.load_kw[day],
            data.actual_pv_kw[day],
            load_forecast.forecast_kw[day],
            data.pv_forecast_hourly_kw[day],
            day_price,
            current_soc,
            update_hours,
            event_triggered,
            previous_pv_end,
            config,
        )
        for target, source in (
            (plan_grid, result.plan_grid),
            (adjusted_grid, result.adjusted_grid),
            (grid_used, result.grid_used),
            (charge, result.charge),
            (discharge, result.discharge),
            (pv_used, result.pv_used),
            (pv_curtailed, result.pv_curtailed),
            (unused_plan, result.unused_plan),
            (emergency, result.emergency),
            (soc, result.soc),
        ):
            target[day] = source
        daily_plan_cost[day] = result.plan_cost_yuan
        daily_adjustment_cost[day] = result.adjustment_cost_yuan
        daily_downward_cost[day] = result.downward_cost_yuan
        daily_upward_cost[day] = result.upward_cost_yuan
        daily_emergency_cost[day] = result.emergency_cost_yuan
        daily_downward_kwh[day] = result.downward_kwh
        daily_upward_kwh[day] = result.upward_kwh
        daily_adjustment_interval_count[day] = result.adjustment_interval_count
        daily_adjustment_event_count[day] = result.adjustment_event_count
        records.extend(result.update_records)
        current_soc = float(result.soc[-1])
        if progress and (day + 1) % 30 == 0:
            print(f"{name}: {day + 1}/{days} days", flush=True)
    return StrategyResult(
        name,
        tuple(update_hours),
        event_triggered,
        plan_grid,
        adjusted_grid,
        grid_used,
        charge,
        discharge,
        pv_used,
        pv_curtailed,
        unused_plan,
        emergency,
        soc,
        daily_plan_cost,
        daily_adjustment_cost,
        daily_downward_cost,
        daily_upward_cost,
        daily_emergency_cost,
        daily_downward_kwh,
        daily_upward_kwh,
        daily_adjustment_interval_count,
        daily_adjustment_event_count,
        records,
    )


def _run_strategy_worker(arguments) -> StrategyResult:
    name, data, load_forecast, updates, event, config = arguments
    return run_strategy(name, data, load_forecast, updates, event, config, progress=False)


def pv_forecast_accuracy(
    data: Problem3Data,
    start_day: int = OUTPUT_START_DAY,
) -> dict[str, dict[str, float | int]]:
    """评价各发布时间在“发布时间至当日24:00”这一实际可执行时域内的精度。"""
    metrics: dict[str, dict[str, float | int]] = {}
    for release_index, release_hour in enumerate(RELEASE_HOURS):
        errors: list[np.ndarray] = []
        for day in range(start_day, len(data.dates)):
            anchor = (
                float(data.actual_pv_kw[day - 1, -1])
                if release_hour == 0 and day > 0
                else (0.0 if release_hour == 0 else float(data.actual_pv_kw[day, release_hour * 6 - 1]))
            )
            forecast = interpolate_pv_release(
                data.pv_forecast_hourly_kw[day, release_index], release_hour, anchor
            )
            actual = data.actual_pv_kw[day, release_hour * 6 :]
            if forecast.shape != actual.shape:
                raise RuntimeError("光伏预报与实际时域未对齐")
            errors.append(forecast - actual)
        merged = np.concatenate(errors)
        metrics[f"{release_hour}:00"] = {
            "sample_count_10min": int(len(merged)),
            "mae_kw": float(np.mean(np.abs(merged))),
            "rmse_kw": float(np.sqrt(np.mean(merged**2))),
            "mean_error_kw": float(np.mean(merged)),
        }
    return metrics


def load_forecast_accuracy(
    data: Problem3Data,
    forecast: LoadForecastResult,
    start_day: int = OUTPUT_START_DAY,
) -> dict[str, dict[str, float | int]]:
    """给出0/6/12/18点负荷预测精度；日内版本只用已执行数据在线校偏。"""
    result: dict[str, dict[str, float | int]] = {}
    for release_hour in RELEASE_HOURS:
        start_slot = release_hour * 6
        errors: list[np.ndarray] = []
        for day in range(start_day, len(data.dates)):
            corrected = online_load_bias_correction(
                forecast.forecast_kw[day], data.load_kw[day], start_slot
            )
            errors.append(corrected[start_slot:] - data.load_kw[day, start_slot:])
        merged = np.concatenate(errors)
        result[f"{release_hour}:00"] = {
            "sample_count_10min": int(len(merged)),
            "mae_kw": float(np.mean(np.abs(merged))),
            "rmse_kw": float(np.sqrt(np.mean(merged**2))),
            "mean_error_kw": float(np.mean(merged)),
        }
    return result


def strategy_metrics(
    result: StrategyResult,
    start_day: int = OUTPUT_START_DAY,
) -> dict[str, float | int | bool]:
    selected = slice(start_day, len(result.plan_grid))
    plan_cost = float(result.daily_plan_cost[selected].sum())
    adjustment_cost = float(result.daily_adjustment_cost[selected].sum())
    emergency_cost = float(result.daily_emergency_cost[selected].sum())
    plan = result.plan_grid[selected]
    adjusted = result.adjusted_grid[selected]
    emergency = result.emergency[selected]
    return {
        "plan_purchase_kwh": float(plan.sum()),
        "final_contracted_purchase_kwh": float(adjusted.sum()),
        "actual_contract_grid_used_kwh": float(result.grid_used[selected].sum()),
        "emergency_purchase_kwh": float(emergency.sum()),
        "total_physical_grid_purchase_kwh": float(result.grid_used[selected].sum() + emergency.sum()),
        "plan_cost_yuan": plan_cost,
        "downward_adjustment_cost_yuan": float(result.daily_downward_cost[selected].sum()),
        "upward_adjustment_cost_yuan": float(result.daily_upward_cost[selected].sum()),
        "adjustment_cost_yuan": adjustment_cost,
        "emergency_cost_yuan": emergency_cost,
        "total_cost_yuan": plan_cost + adjustment_cost + emergency_cost,
        "downward_adjustment_kwh": float(result.daily_downward_kwh[selected].sum()),
        "upward_adjustment_kwh": float(result.daily_upward_kwh[selected].sum()),
        "total_adjustment_kwh": float(
            result.daily_downward_kwh[selected].sum() + result.daily_upward_kwh[selected].sum()
        ),
        "adjustment_interval_count": int(result.daily_adjustment_interval_count[selected].sum()),
        "adjustment_event_count": int(result.daily_adjustment_event_count[selected].sum()),
        "emergency_interval_count": int(np.sum(emergency > TOLERANCE)),
        "emergency_day_count": int(np.sum(emergency.sum(axis=1) > TOLERANCE)),
        "unused_contracted_energy_kwh": float(result.unused_plan[selected].sum()),
        "charge_kwh": float(result.charge[selected].sum()),
        "discharge_kwh": float(result.discharge[selected].sum()),
        "pv_consumed_kwh": float(result.pv_used[selected].sum()),
        "pv_curtailed_kwh": float(result.pv_curtailed[selected].sum()),
        "event_triggered": result.event_triggered,
    }


def marginal_information_value(
    metrics: dict[str, dict[str, float | int | bool]],
) -> dict[str, dict[str, float | bool]]:
    pairs = (("6:00", "S0", "S1"), ("12:00", "S1", "S2"), ("18:00", "S2", "S3"))
    values: dict[str, dict[str, float | bool]] = {}
    for release, before, after in pairs:
        saving = float(metrics[before]["total_cost_yuan"]) - float(metrics[after]["total_cost_yuan"])
        values[f"delta_C_{release.split(':')[0]}"] = {
            "before_cost_yuan": float(metrics[before]["total_cost_yuan"]),
            "after_cost_yuan": float(metrics[after]["total_cost_yuan"]),
            "marginal_value_yuan": saving,
            "reduces_actual_total_cost": bool(saving > 0.0),
        }
    return values


def audit_problem3(
    result: StrategyResult,
    data: Problem3Data,
    load_forecast: LoadForecastResult,
    config: DispatchConfig = DispatchConfig(),
) -> dict[str, object]:
    load_kwh = kw_to_interval_kwh(data.load_kw, config.interval_hours)
    physical = audit_dispatch(
        result.grid_used,
        result.charge,
        result.discharge,
        result.pv_used,
        result.emergency,
        load_kwh,
        result.soc,
        config,
    )
    continuity_error = float(np.max(np.abs(result.soc[:-1, -1] - result.soc[1:, 0])))
    inherited_errors = [
        abs(record.inherited_actual_soc_kwh - result.soc[record.day_index, record.start_slot])
        for record in result.update_records
    ]
    contract_identity = result.adjusted_grid - result.grid_used - result.unused_plan
    pv_partition = kw_to_interval_kwh(data.actual_pv_kw, config.interval_hours) - result.pv_used - result.pv_curtailed
    finite_arrays = (
        result.plan_grid,
        result.adjusted_grid,
        result.grid_used,
        result.charge,
        result.discharge,
        result.pv_used,
        result.pv_curtailed,
        result.unused_plan,
        result.emergency,
        result.soc,
    )
    result_audit: dict[str, object] = dict(physical)
    result_audit.update(
        {
            "period_count_per_day": int(result.plan_grid.shape[1]),
            "all_365_days_present": bool(result.plan_grid.shape == (365, 144)),
            "ten_minute_energy_factor_hours": config.interval_hours,
            "interval_power_limit_kwh": config.interval_limit_kwh,
            "price_source": "附件1.xlsx",
            "actual_source": "附件2.xlsx",
            "pv_forecast_source": "附件3.xlsx",
            "attachment4_opened": False,
            "no_reverse_sale": bool(
                result.plan_grid.min() >= -TOLERANCE
                and result.adjusted_grid.min() >= -TOLERANCE
                and result.grid_used.min() >= -TOLERANCE
            ),
            "max_contract_partition_error_kwh": float(np.max(np.abs(contract_identity))),
            "max_pv_partition_error_kwh": float(np.max(np.abs(pv_partition))),
            "soc_day_continuity_error_kwh": continuity_error,
            "max_update_soc_inheritance_error_kwh": float(max(inherited_errors, default=0.0)),
            "executed_intervals_frozen": True,
            "load_forecast_leakage_detected": load_forecast.leakage_detected,
            "pv_future_actual_used_in_optimization": False,
            "all_arrays_finite": bool(all(np.all(np.isfinite(array)) for array in finite_arrays)),
            "time_alignment": (
                "附件2的0:10对应0:00-0:10、0:00(+1)对应23:50-24:00；"
                "附件3预测1小时对应发布后第1个整点，线性插值到区间终点"
            ),
            "daily_planned_terminal_policy": "每次优化以当日0:00实际SOC为24:00目标；实际偏差由下一日当前SOC继承",
        }
    )
    required_true = (
        "all_365_days_present",
        "no_reverse_sale",
        "executed_intervals_frozen",
        "all_arrays_finite",
        "soc_bounds_ok",
        "power_limits_ok",
        "no_simultaneous_charge_discharge",
    )
    numeric_limits = (
        abs(float(result_audit["max_power_balance_error_kwh"])) <= TOLERANCE,
        abs(float(result_audit["max_soc_dynamics_error_kwh"])) <= TOLERANCE,
        abs(float(result_audit["max_contract_partition_error_kwh"])) <= TOLERANCE,
        abs(float(result_audit["max_pv_partition_error_kwh"])) <= TOLERANCE,
        continuity_error <= TOLERANCE,
        float(result_audit["max_update_soc_inheritance_error_kwh"]) <= TOLERANCE,
    )
    result_audit["all_required_checks_passed"] = bool(
        all(bool(result_audit[key]) for key in required_true)
        and not load_forecast.leakage_detected
        and all(numeric_limits)
    )
    return result_audit


def run_problem3_study(
    data: Problem3Data,
    config: DispatchConfig = DispatchConfig(),
    include_event_trigger: bool = True,
    progress: bool = True,
    parallel: bool = True,
) -> Problem3Study:
    load_forecast = build_causal_load_forecasts(data.load_kw)
    definitions: list[tuple[str, tuple[int, ...], bool]] = [
        ("S0", (), False),
        ("S1", (6,), False),
        ("S2", (6, 12), False),
        ("S3", (6, 12, 18), False),
    ]
    if include_event_trigger:
        definitions.append(("S3E", (6, 12, 18), True))
    strategies: dict[str, StrategyResult] = {}
    metrics: dict[str, dict[str, float | int | bool]] = {}
    if parallel and len(definitions) > 1:
        from concurrent.futures import ProcessPoolExecutor, as_completed

        worker_count = min(4, len(definitions))
        arguments = [
            (name, data, load_forecast, updates, event, config)
            for name, updates, event in definitions
        ]
        with ProcessPoolExecutor(max_workers=worker_count) as executor:
            futures = {executor.submit(_run_strategy_worker, item): item[0] for item in arguments}
            for future in as_completed(futures):
                name = futures[future]
                result = future.result()
                strategies[name] = result
                if progress:
                    print(f"{name}: annual simulation complete", flush=True)
    else:
        for name, updates, event in definitions:
            strategies[name] = run_strategy(name, data, load_forecast, updates, event, config, progress)
    # 保持论文表格的固定策略顺序，不受并行任务完成先后影响。
    strategies = {name: strategies[name] for name, _, _ in definitions}
    for name, result in strategies.items():
        metrics[name] = strategy_metrics(result)
    marginal = marginal_information_value(metrics)
    selected_name = min(metrics, key=lambda key: (float(metrics[key]["total_cost_yuan"]), key))
    selected = strategies[selected_name]
    audit = audit_problem3(selected, data, load_forecast, config)
    return Problem3Study(
        selected,
        strategies,
        metrics,
        pv_forecast_accuracy(data),
        marginal,
        load_forecast,
        audit,
    )


def _rolling_result(result: StrategyResult) -> RollingYearResult:
    return RollingYearResult(
        result.plan_grid,
        result.adjusted_grid,
        result.charge,
        result.discharge,
        result.pv_used,
        result.emergency,
        result.soc,
        result.daily_plan_cost,
        result.daily_adjustment_cost,
        result.daily_emergency_cost,
        result.daily_downward_kwh,
        result.daily_upward_kwh,
    )


def write_result3_workbook(
    root: str | Path,
    output_path: str | Path,
    study: Problem3Study,
    price: np.ndarray | Sequence[float],
) -> Path:
    """严格复制result3模板并写入最小实际总成本策略，不改工作表名/表头/格式。"""
    static_price = np.asarray(price, dtype=float).reshape(-1)
    if static_price.shape != (144,):
        raise ValueError("写入result3时附件1电价必须为144点")
    price_matrix = np.repeat(static_price[None, :], 365, axis=0)
    return write_rolling_workbook(
        root,
        output_path,
        "result3.xlsx",
        _rolling_result(study.selected_result),
        price_matrix,
    )


def _jsonable(value):
    if isinstance(value, dict):
        return {str(key): _jsonable(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_jsonable(item) for item in value]
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, (np.integer, np.floating)):
        return value.item()
    if isinstance(value, np.bool_):
        return bool(value)
    return value


def write_analysis_outputs(
    root: str | Path,
    data: Problem3Data,
    study: Problem3Study,
) -> tuple[Path, Path, Path, Path, Path]:
    import csv

    root_path = Path(root).resolve()
    json_path = root_path / "problem3_analysis.json"
    report_path = root_path / "problem3_analysis_report.md"
    strategy_csv = root_path / "problem3_strategy_comparison.csv"
    pv_csv = root_path / "problem3_pv_accuracy.csv"
    update_csv = root_path / "problem3_update_log.csv"
    load_accuracy = load_forecast_accuracy(data, study.load_forecast)
    selected_counts = {
        name: int(np.sum(study.load_forecast.selected_config[OUTPUT_START_DAY:] == name))
        for name in sorted(set(study.load_forecast.selected_config[OUTPUT_START_DAY:]))
    }
    payload = {
        "input_whitelist": list(data.source_files),
        "forbidden_input": "附件4.xlsx（未打开）",
        "time_alignment": study.audit["time_alignment"],
        "units": {"power": "kW", "interval": "10 min", "energy_factor_h": 1.0 / 6.0, "energy": "kWh"},
        "load_prediction": {
            "windows_days": list(LOAD_WINDOWS),
            "decays": list(LOAD_DECAYS),
            "rolling_validation_days": VALIDATION_DAYS,
            "selected_config_counts_2025_02_01_to_12_31": selected_counts,
            "accuracy_by_release": load_accuracy,
            "future_leakage_detected": study.load_forecast.leakage_detected,
        },
        "pv_accuracy_by_release": study.pv_accuracy,
        "strategy_comparison": study.strategy_metrics,
        "marginal_information_value": study.marginal_information_value,
        "selected_strategy_for_result3": study.selected_result.name,
        "audit": study.audit,
    }
    json_path.write_text(json.dumps(_jsonable(payload), ensure_ascii=False, indent=2), encoding="utf-8")

    strategy_fields = ["strategy", *next(iter(study.strategy_metrics.values())).keys()]
    with strategy_csv.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=strategy_fields)
        writer.writeheader()
        for name, values in study.strategy_metrics.items():
            writer.writerow({"strategy": name, **values})
    with pv_csv.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=["release", "sample_count_10min", "mae_kw", "rmse_kw", "mean_error_kw"])
        writer.writeheader()
        for release, values in study.pv_accuracy.items():
            writer.writerow({"release": release, **values})
    update_fields = ["strategy", *UpdateRecord.__dataclass_fields__.keys()]
    with update_csv.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=update_fields)
        writer.writeheader()
        for name, strategy in study.strategies.items():
            for record in strategy.update_records:
                writer.writerow({"strategy": name, **asdict(record)})

    lines = [
        "# 问题三：多阶段MPC滚动修正分析",
        "",
        f"最终写入 result3.xlsx 的策略：**{study.selected_result.name}**。",
        "",
        "时间对齐：附件2每列是10分钟区间终点；附件3的预测1小时是发布后第1个整点。"
        "0/6/12/18点先以当时已观测光伏为锚点线性插值，再只优化尚未执行的后缀。",
        "",
        "## 负荷预测滚动选择",
        "",
        "候选窗口固定为7/14/21/28日，候选衰减参数为0.65/0.80/0.90/1.00；"
        "每日参数只依据此前28个已结束日期的滚动RMSE选择。",
        "",
        "| 选中配置 | 天数 |",
        "|---|---:|",
    ]
    for config_name, count in selected_counts.items():
        lines.append(f"| {config_name} | {count} |")
    lines.extend(
        [
            "",
            "| 发布时间 | 负荷MAE/kW | 负荷RMSE/kW |",
            "|---|---:|---:|",
        ]
    )
    for release, values in load_accuracy.items():
        lines.append(f"| {release} | {values['mae_kw']:.3f} | {values['rmse_kw']:.3f} |")
    lines.extend(
        [
        "",
        "## 光伏预测精度",
        "",
        "| 发布时间 | 样本数 | MAE/kW | RMSE/kW | 平均误差/kW |",
        "|---|---:|---:|---:|---:|",
        ]
    )
    for release, values in study.pv_accuracy.items():
        lines.append(
            f"| {release} | {values['sample_count_10min']} | {values['mae_kw']:.3f} | "
            f"{values['rmse_kw']:.3f} | {values['mean_error_kw']:.3f} |"
        )
    lines.extend(
        [
            "",
            "## 策略成本对比（2025-02-01至2025-12-31）",
            "",
            "| 策略 | 总费用/元 | 计划费/元 | 调整费/元 | 紧急购电费/元 | 紧急电量/kWh | 调整电量/kWh | 调整次数 |",
            "|---|---:|---:|---:|---:|---:|---:|---:|",
        ]
    )
    for name, values in study.strategy_metrics.items():
        lines.append(
            f"| {name} | {values['total_cost_yuan']:.2f} | {values['plan_cost_yuan']:.2f} | "
            f"{values['adjustment_cost_yuan']:.2f} | {values['emergency_cost_yuan']:.2f} | "
            f"{values['emergency_purchase_kwh']:.2f} | {values['total_adjustment_kwh']:.2f} | "
            f"{values['adjustment_event_count']} |"
        )
    lines.extend(["", "## 新增预测的边际信息价值", "", "| 新增时点 | 边际价值/元 | 是否降低总成本 |", "|---|---:|---|"])
    for name, values in study.marginal_information_value.items():
        lines.append(
            f"| {name.replace('delta_C_', '')}:00 | {values['marginal_value_yuan']:.2f} | "
            f"{'是' if values['reduces_actual_total_cost'] else '否'} |"
        )
    lines.extend(
        [
            "",
            "## 关键约束审计",
            "",
            f"- 最大电量平衡误差：{study.audit['max_power_balance_error_kwh']:.3e} kWh。",
            f"- 最大SOC递推误差：{study.audit['max_soc_dynamics_error_kwh']:.3e} kWh。",
            f"- SOC范围：[{study.audit['soc_min_kwh']:.3f}, {study.audit['soc_max_kwh']:.3f}] kWh。",
            f"- 最大10分钟充/放电量：{study.audit['max_charge_interval_kwh']:.3f}/"
            f"{study.audit['max_discharge_interval_kwh']:.3f} kWh。",
            f"- 当前实际SOC继承最大误差：{study.audit['max_update_soc_inheritance_error_kwh']:.3e} kWh。",
            f"- 全部检查通过：{study.audit['all_required_checks_passed']}。",
        ]
    )
    report_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return json_path, report_path, strategy_csv, pv_csv, update_csv


def print_key_results(study: Problem3Study) -> None:
    print("\n问题三策略对比（2025-02-01至2025-12-31）")
    for name, values in study.strategy_metrics.items():
        print(
            f"{name}: total={values['total_cost_yuan']:.2f} yuan, "
            f"plan={values['plan_cost_yuan']:.2f}, adjust={values['adjustment_cost_yuan']:.2f}, "
            f"emergency={values['emergency_cost_yuan']:.2f}, "
            f"emergency_energy={values['emergency_purchase_kwh']:.2f} kWh, "
            f"adjust_energy={values['total_adjustment_kwh']:.2f} kWh"
        )
    print("\n新增预测边际价值")
    for name, values in study.marginal_information_value.items():
        print(f"{name}: {values['marginal_value_yuan']:.2f} yuan")
    print(f"\nresult3 selected strategy: {study.selected_result.name}")
    print(json.dumps(_jsonable(study.audit), ensure_ascii=False, indent=2))


def main() -> None:
    import argparse

    parser = argparse.ArgumentParser(description="问题三：多阶段滚动预测 + MPC + 非对称调整成本")
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parent)
    parser.add_argument("--output", type=Path, default=None)
    parser.add_argument("--no-event-trigger", action="store_true", help="不运行额外的S3E事件触发策略")
    parser.add_argument("--quiet", action="store_true")
    args = parser.parse_args()
    root = args.root.resolve()
    output = (args.output or (root / "result3.xlsx")).resolve()
    data = load_problem3_data(root, print_structure=not args.quiet)
    study = run_problem3_study(
        data,
        include_event_trigger=not args.no_event_trigger,
        progress=not args.quiet,
    )
    write_result3_workbook(root, output, study, data.static_price)
    write_analysis_outputs(root, data, study)
    print_key_results(study)


if __name__ == "__main__":
    main()

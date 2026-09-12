import math
import tempfile
import unittest
from pathlib import Path

import numpy as np

from microgrid_optimization import (
    DispatchConfig,
    DeterministicYearResult,
    RollingYearResult,
    adjustment_settlement,
    audit_dispatch,
    build_time_labels,
    expand_hourly_forecast,
    evaluate_midnight_only_year,
    group_emergency_intervals,
    kw_to_interval_kwh,
    load_input_data,
    run_rolling_day,
    soc_path,
    solve_dispatch,
    write_problem1_workbook,
    write_deterministic_workbook,
    write_rolling_workbook,
)
from problem2_optimization import (
    ALPHA_CANDIDATES,
    build_causal_forecasts,
    calculate_problem2_cost,
    optimize_day_ahead,
    simulate_realtime,
)


class TimeAndUnitTests(unittest.TestCase):
    def test_kw_is_converted_to_ten_minute_kwh(self):
        result = kw_to_interval_kwh(np.array([600.0, 1200.0]))
        np.testing.assert_allclose(result, [100.0, 200.0])

    def test_hourly_forecast_is_held_for_six_intervals(self):
        result = expand_hourly_forecast(np.array([10.0, 20.0]))
        np.testing.assert_allclose(result, [10.0] * 6 + [20.0] * 6)

    def test_source_end_timestamps_map_to_same_day_physical_intervals(self):
        labels = build_time_labels()
        self.assertEqual(len(labels), 144)
        self.assertEqual(labels[0], "0:00-0:10")
        self.assertEqual(labels[5], "0:50-1:00")
        self.assertEqual(labels[6], "1:00-1:10")
        self.assertEqual(labels[-1], "23:50-0:00+1")

    def test_original_attachments_are_loaded_into_expected_tensors(self):
        root = Path(__file__).resolve().parents[1]
        data = load_input_data(root)
        self.assertEqual(data.static_price.shape, (144,))
        self.assertEqual(data.load_kw.shape, (365, 144))
        self.assertEqual(data.actual_pv_kw.shape, (365, 144))
        self.assertEqual(data.forecast_pv_kw.shape, (365, 4, 24))
        self.assertEqual(data.dynamic_price.shape, (365, 144))
        self.assertEqual(str(data.dates[0]), "2025-01-01")
        self.assertEqual(str(data.dates[-1]), "2025-12-31")
        self.assertTrue(np.all(np.isfinite(data.load_kw)))
        self.assertTrue(np.all(data.dynamic_price >= 0.0))

    def test_midnight_only_baseline_returns_finite_nonnegative_metrics(self):
        root = Path(__file__).resolve().parents[1]
        data = load_input_data(root)
        prices = np.broadcast_to(data.static_price, (365, 144))
        baseline = evaluate_midnight_only_year(data, prices)
        self.assertGreater(baseline["plan_purchase_kwh"], 0.0)
        self.assertGreaterEqual(baseline["emergency_purchase_kwh"], 0.0)
        self.assertTrue(math.isfinite(baseline["total_cost_yuan"]))


class StorageAndOptimizationTests(unittest.TestCase):
    def test_soc_uses_ninety_percent_charge_and_discharge_efficiency(self):
        result = soc_path(6000.0, np.array([100.0, 0.0]), np.array([0.0, 90.0]), 0.9)
        np.testing.assert_allclose(result, [6000.0, 6090.0, 5990.0])

    def test_dispatch_shifts_energy_from_cheap_to_expensive_interval(self):
        config = DispatchConfig(power_limit_kw=5000.0, throughput_penalty=1e-8)
        result = solve_dispatch(
            load_kwh=np.array([100.0, 100.0]),
            pv_kwh=np.zeros(2),
            price=np.array([1.0, 3.0]),
            initial_soc=6000.0,
            terminal_soc=6000.0,
            config=config,
        )
        self.assertTrue(result.success, result.message)
        np.testing.assert_allclose(
            result.grid + result.pv_used + result.discharge,
            np.array([100.0, 100.0]) + result.charge,
            atol=1e-6,
        )
        self.assertLess(result.grid[1], 1e-6)
        self.assertGreater(result.charge[0], 100.0)
        self.assertLess(np.max(np.minimum(result.charge, result.discharge)), 1e-6)
        self.assertAlmostEqual(result.soc[0], 6000.0, places=6)
        self.assertAlmostEqual(result.soc[-1], 6000.0, places=6)

    def test_dispatch_audit_reports_small_residuals_and_valid_bounds(self):
        load = np.array([100.0, 100.0])
        pv = np.zeros(2)
        result = solve_dispatch(load, pv, np.array([1.0, 3.0]), 6000.0, 6000.0)
        audit = audit_dispatch(result.grid, result.charge, result.discharge, result.pv_used, np.zeros(2), load, result.soc)
        self.assertLess(audit["max_power_balance_error_kwh"], 1e-6)
        self.assertLess(audit["max_soc_dynamics_error_kwh"], 1e-6)
        self.assertTrue(audit["soc_bounds_ok"])
        self.assertTrue(audit["power_limits_ok"])
        self.assertTrue(audit["no_simultaneous_charge_discharge"])

    def test_rolling_day_has_no_adjustment_or_emergency_under_perfect_forecast(self):
        load = np.full(144, 100.0)
        actual_pv = np.zeros(144)
        hourly_forecasts = np.zeros((4, 24))
        price = np.ones(144)
        result = run_rolling_day(
            load_kwh=load,
            actual_pv_kwh=actual_pv,
            hourly_forecast_kw=hourly_forecasts,
            price=price,
            initial_soc=6000.0,
            terminal_soc=6000.0,
        )
        np.testing.assert_allclose(result.final_grid, result.plan.grid, atol=1e-6)
        np.testing.assert_allclose(result.emergency, 0.0, atol=1e-6)
        self.assertAlmostEqual(result.adjustment.penalty_cost, 0.0, places=6)
        self.assertAlmostEqual(result.adjustment.incremental_cost, 0.0, places=6)
        self.assertAlmostEqual(result.total_cost, result.plan.purchase_cost, places=5)
        self.assertAlmostEqual(result.soc[0], 6000.0, places=6)
        self.assertAlmostEqual(result.soc[-1], 6000.0, places=6)

    def test_rolling_adjustment_never_cycles_battery_to_avoid_adjustment_fees(self):
        root = Path(__file__).resolve().parents[1]
        data = load_input_data(root)
        day = 1  # 2025-01-02 reproduced the original LP degeneracy at slot 37.
        result = run_rolling_day(
            kw_to_interval_kwh(data.load_kw[day]),
            kw_to_interval_kwh(data.actual_pv_kw[day]),
            data.forecast_pv_kw[day],
            data.static_price,
            initial_soc=6000.0,
            terminal_soc=6000.0,
        )
        self.assertLess(np.max(np.minimum(result.charge, result.discharge)), 1e-6)


class CausalProblem2Tests(unittest.TestCase):
    def test_day_forecast_is_unchanged_when_current_and_future_actuals_change(self):
        rng = np.random.default_rng(2026)
        actual = rng.normal(500.0, 80.0, size=(20, 144))
        changed = actual.copy()
        changed[10:] += 10000.0

        original = build_causal_forecasts(actual, np.ones(144))
        perturbed = build_causal_forecasts(changed, np.ones(144))

        np.testing.assert_allclose(original.mean[10], perturbed.mean[10])
        np.testing.assert_allclose(original.risk_adjusted[10], perturbed.risk_adjusted[10])
        self.assertEqual(original.selected_alpha[10], perturbed.selected_alpha[10])
        self.assertEqual(original.max_history_day_used[10], 9)
        self.assertFalse(original.leakage_detected)

    def test_cold_start_and_risk_alpha_are_explicit_and_causal(self):
        actual = np.full((9, 144), 100.0)
        forecasts = build_causal_forecasts(actual, np.ones(144))
        np.testing.assert_allclose(forecasts.mean[0], 0.0)
        self.assertAlmostEqual(forecasts.selected_alpha[0], 0.80)
        self.assertTrue(set(np.unique(forecasts.selected_alpha)).issubset(set(ALPHA_CANDIDATES)))

    def test_realtime_simulation_prioritizes_storage_then_emergency(self):
        config = DispatchConfig()
        plan = np.array([0.0, 1000.0, 0.0])
        load = np.array([100.0, 0.0, 1000.0])
        pv = np.zeros(3)
        result = simulate_realtime(plan, load, pv, initial_soc=1200.0, config=config)

        self.assertAlmostEqual(result.emergency[0], 100.0)
        self.assertAlmostEqual(result.charge[1], config.interval_limit_kwh)
        self.assertAlmostEqual(result.discharge[2], 675.0, places=6)
        self.assertAlmostEqual(result.emergency[2], 325.0, places=6)
        self.assertAlmostEqual(result.soc[0], 1200.0)
        self.assertAlmostEqual(result.soc[-1], 1200.0)
        np.testing.assert_allclose(
            result.grid_used + pv - result.pv_curtailed + result.discharge + result.emergency,
            load + result.charge,
            atol=1e-8,
        )
        np.testing.assert_allclose(result.grid_used + result.unused_plan, plan, atol=1e-8)
        self.assertLessEqual(result.charge.max(), config.interval_limit_kwh + 1e-8)
        self.assertLessEqual(result.discharge.max(), config.interval_limit_kwh + 1e-8)

    def test_realtime_uses_pv_before_plan_and_never_exports(self):
        config = DispatchConfig()
        plan = np.array([100.0, 0.0])
        load = np.array([50.0, 10000.0])
        pv = np.array([80.0, 0.0])
        result = simulate_realtime(plan, load, pv, initial_soc=10800.0, config=config)

        self.assertAlmostEqual(result.grid_used[0], 0.0)
        self.assertAlmostEqual(result.unused_plan[0], 100.0)
        self.assertAlmostEqual(result.pv_curtailed[0], 30.0)
        self.assertAlmostEqual(result.emergency[0], 0.0)
        self.assertAlmostEqual(
            result.emergency[1], 10000.0 - config.interval_limit_kwh
        )
        self.assertTrue(np.all(result.grid_used >= -1e-8))
        self.assertTrue(np.all(result.pv_curtailed <= pv + 1e-8))

    def test_day_ahead_optimizer_uses_only_forecast_and_current_soc(self):
        forecast_net = np.full(144, 100.0)
        plan = optimize_day_ahead(forecast_net, np.ones(144), initial_soc=6000.0)
        self.assertTrue(plan.success, plan.message)
        self.assertEqual(plan.grid.shape, (144,))
        self.assertGreaterEqual(plan.grid.min(), -1e-8)
        self.assertAlmostEqual(plan.soc[0], 6000.0, places=6)
        self.assertAlmostEqual(plan.soc[-1], 6000.0, places=6)

    def test_problem2_cost_charges_frozen_plan_and_five_times_emergency(self):
        plan = np.array([10.0, 20.0])
        emergency = np.array([3.0, 4.0])
        price = np.array([2.0, 5.0])
        cost = calculate_problem2_cost(plan, emergency, price)
        self.assertAlmostEqual(cost.normal_cost, 120.0)
        self.assertAlmostEqual(cost.emergency_cost, 130.0)
        self.assertAlmostEqual(cost.total_cost, 250.0)


class SettlementAndReportingTests(unittest.TestCase):
    def test_adjustment_cost_uses_fifty_and_one_hundred_fifty_percent_rates(self):
        plan = np.array([100.0, 100.0])
        adjusted = np.array([80.0, 130.0])
        price = np.array([2.0, 2.0])
        result = adjustment_settlement(plan, adjusted, price)
        self.assertAlmostEqual(result.downward_kwh, 20.0)
        self.assertAlmostEqual(result.upward_kwh, 30.0)
        self.assertAlmostEqual(result.penalty_cost, 20.0)
        self.assertAlmostEqual(result.incremental_cost, 90.0)

    def test_emergency_intervals_are_grouped_when_adjacent(self):
        amounts = np.array([0.0, 10.0, 20.0, 0.0, 5.0])
        groups = group_emergency_intervals(amounts)
        self.assertEqual(groups, [("0:10-0:30", 30.0), ("0:40-0:50", 5.0)])

    def test_problem1_writer_corrects_time_labels_and_preserves_styles(self):
        from openpyxl import load_workbook

        root = Path(__file__).resolve().parents[1]
        data = load_input_data(root)
        dispatch = solve_dispatch(
            kw_to_interval_kwh(data.problem1_load_kw),
            kw_to_interval_kwh(data.problem1_pv_kw),
            data.static_price,
            initial_soc=6000.0,
            terminal_soc=6000.0,
        )
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "result1.xlsx"
            write_problem1_workbook(root, output, dispatch)
            template = load_workbook(root / "附件" / "附件5" / "result1.xlsx")
            written = load_workbook(output, data_only=False)
            self.assertEqual(written.sheetnames, template.sheetnames)
            for sheet_name in template.sheetnames:
                self.assertEqual(written[sheet_name]["A1"].value, template[sheet_name]["A1"].value)
                self.assertEqual(written[sheet_name]["A1"].style_id, template[sheet_name]["A1"].style_id)
            self.assertEqual(written["计划购电量"]["A2"].value, "0:00-0:10")
            self.assertEqual(written["计划购电量"]["A145"].value, "23:50-0:00+1")
            self.assertEqual(
                written["计划购电量"]["A2"].style_id,
                template["计划购电量"]["A2"].style_id,
            )
            self.assertIsInstance(written["计划购电量"]["B2"].value, float)
            self.assertAlmostEqual(written["充放电量"]["E2"].value, 6000.0)
            self.assertAlmostEqual(written["充放电量"]["E3"].value, 6000.0)

    def test_annual_writer_expands_template_with_aligned_time_headers(self):
        from openpyxl import load_workbook

        root = Path(__file__).resolve().parents[1]
        zeros = np.zeros((365, 144))
        soc = np.full((365, 145), 6000.0)
        result = DeterministicYearResult(zeros, zeros, zeros, zeros, zeros, soc, np.zeros(365), 0.0)
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "result2.xlsx"
            write_deterministic_workbook(root, output, "result2.xlsx", result, np.ones((365, 144)))
            workbook = load_workbook(output, data_only=False)
            self.assertEqual(workbook.sheetnames, ["计划购电量", "充放电量", "紧急购电量"])
            self.assertEqual(workbook["计划购电量"]["B1"].value, "0:00-0:10")
            self.assertEqual(workbook["计划购电量"]["EO1"].value, "23:50-0:00+1")
            self.assertEqual(workbook["计划购电量"]["EQ1"].value, "全天购电费")
            self.assertEqual(workbook["充放电量"].max_row, 2005)
            self.assertEqual(workbook["紧急购电量"].max_row, 1003)
            self.assertEqual(workbook["充放电量"]["A2"].value.strftime("%Y-%m-%d"), "2025-02-01")
            self.assertEqual(workbook["充放电量"]["A2000"].value.strftime("%Y-%m-%d"), "2025-12-31")

    def test_rolling_writer_populates_plan_and_adjusted_sheets(self):
        from openpyxl import load_workbook

        root = Path(__file__).resolve().parents[1]
        plan = np.full((365, 144), 10.0)
        adjusted = np.full((365, 144), 12.0)
        zeros = np.zeros((365, 144))
        soc = np.full((365, 145), 6000.0)
        daily_plan_cost = np.full(365, 1440.0)
        daily_adjustment_cost = np.full(365, 432.0)
        result = RollingYearResult(
            plan, adjusted, zeros, zeros, zeros, zeros, soc,
            daily_plan_cost, daily_adjustment_cost, np.zeros(365), np.zeros(365), np.full(365, 288.0),
        )
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "result3.xlsx"
            write_rolling_workbook(root, output, "result3.xlsx", result, np.ones((365, 144)))
            workbook = load_workbook(output, data_only=False)
            self.assertEqual(workbook.sheetnames, ["计划购电量", "调整购电量", "充放电量", "紧急购电量"])
            self.assertAlmostEqual(workbook["计划购电量"]["B2"].value, 10.0)
            self.assertAlmostEqual(workbook["调整购电量"]["B2"].value, 12.0)
            self.assertAlmostEqual(workbook["调整购电量"]["EQ2"].value, 1872.0)


if __name__ == "__main__":
    unittest.main()

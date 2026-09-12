"""问题四：在问题二因果预测和问题三滚动MPC基础上接入附件4动态电价。"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Mapping, Sequence

import numpy as np

from microgrid_optimization import (
    DispatchConfig,
    RollingYearResult,
    write_rolling_workbook,
)
from problem2_optimization import (
    Problem2Data,
    Problem2ModelStudy,
    Problem2YearResult,
    audit_problem2,
    run_model_study,
    save_result2,
)
from problem3_optimization import (
    LoadForecastResult,
    OUTPUT_START_DAY,
    RELEASE_HOURS,
    Problem3Data,
    StrategyResult,
    _find_sheet,
    _parse_date,
    audit_problem3,
    build_causal_load_forecasts,
    interpolate_pv_release,
    run_strategy,
    strategy_metrics,
)


@dataclass(frozen=True)
class Problem4Data:
    dates: np.ndarray
    load_kw: np.ndarray
    actual_pv_kw: np.ndarray
    pv_forecast_hourly_kw: np.ndarray
    dynamic_price: np.ndarray
    source_files: tuple[str, str, str]

    def as_problem2_data(self, price: np.ndarray | None = None) -> Problem2Data:
        return Problem2Data(
            self.dynamic_price if price is None else np.asarray(price, dtype=float),
            self.dates,
            self.load_kw,
            self.actual_pv_kw,
        )

    def as_problem3_data(self, price: np.ndarray | None = None) -> Problem3Data:
        return Problem3Data(
            self.dynamic_price if price is None else np.asarray(price, dtype=float),
            self.dates,
            self.load_kw,
            self.actual_pv_kw,
            self.pv_forecast_hourly_kw,
            self.source_files,
        )


@dataclass(frozen=True)
class PVResidualAnalysis:
    statistics: dict[str, dict[str, float | int]]
    residuals_kw: dict[str, np.ndarray]
    acf: dict[str, np.ndarray]
    error_decreases_with_shorter_horizon: bool


@dataclass(frozen=True)
class Problem4Study:
    price_blind_study: Problem2ModelStudy
    dynamic_problem2_study: Problem2ModelStudy
    dynamic_mpc_result: StrategyResult
    load_forecast: LoadForecastResult
    residual_analysis: PVResidualAnalysis
    strategy_comparison: dict[str, dict[str, float | int | bool]]
    arbitrage_analysis: dict[str, dict[str, float]]
    audit_problem4_2: dict[str, object]
    audit_problem4_3: dict[str, object]


def problem4_input_files(root: str | Path) -> tuple[Path, Path, Path]:
    """问题四输入白名单；不读取附件1或现有结果文件。"""
    folder = Path(root).resolve() / "附件"
    paths = tuple(folder / name for name in ("附件2.xlsx", "附件3.xlsx", "附件4.xlsx"))
    missing = [path.name for path in paths if not path.is_file()]
    if missing:
        raise FileNotFoundError(f"问题四缺少输入文件: {missing}")
    return paths  # type: ignore[return-value]


def _describe_workbook(path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    print(f"{path.name}: sheets={workbook.sheetnames}")
    for sheet in workbook.worksheets:
        print(f"  {sheet.title}: shape=({sheet.max_row}, {sheet.max_column})")
    workbook.close()


def load_problem4_data(root: str | Path, print_structure: bool = True) -> Problem4Data:
    """读取附件2实际量、附件3光伏预测和附件4全年动态电价。"""
    from openpyxl import load_workbook

    actual_path, forecast_path, price_path = problem4_input_files(root)
    if print_structure:
        for path in (actual_path, forecast_path, price_path):
            _describe_workbook(path)

    workbook2 = load_workbook(actual_path, read_only=True, data_only=True)
    load_sheet = _find_sheet(workbook2, "负载")
    pv_sheet = _find_sheet(workbook2, "实际")
    load_rows = list(load_sheet.iter_rows(min_row=2, values_only=True))
    pv_rows = list(pv_sheet.iter_rows(min_row=2, values_only=True))
    workbook2.close()
    dates = np.asarray([_parse_date(row[0]) for row in load_rows], dtype="datetime64[D]")
    pv_dates = np.asarray([_parse_date(row[0]) for row in pv_rows], dtype="datetime64[D]")
    load_kw = np.asarray([row[1:] for row in load_rows], dtype=float)
    actual_pv_kw = np.asarray([row[1:] for row in pv_rows], dtype=float)

    workbook3 = load_workbook(forecast_path, read_only=True, data_only=True)
    forecast_rows = list(workbook3.worksheets[0].iter_rows(min_row=2, values_only=True))
    workbook3.close()
    if len(forecast_rows) != 365 * 4:
        raise ValueError(f"附件3应为365×4条预报，实际为{len(forecast_rows)}")
    forecast_dates: list[np.datetime64] = []
    issue_hours: list[int] = []
    forecasts: list[list[float]] = []
    current_date: np.datetime64 | None = None
    for row in forecast_rows:
        if row[0] not in (None, ""):
            current_date = _parse_date(row[0])
        if current_date is None:
            raise ValueError("附件3首条记录缺少日期")
        forecast_dates.append(current_date)
        issue_hours.append(int(str(row[1]).split(":")[0]))
        forecasts.append([float(value) for value in row[2:26]])
    forecast_array = np.asarray(forecasts, dtype=float).reshape(365, 4, 24)
    forecast_date_matrix = np.asarray(forecast_dates, dtype="datetime64[D]").reshape(365, 4)
    issue_matrix = np.asarray(issue_hours, dtype=int).reshape(365, 4)

    workbook4 = load_workbook(price_path, read_only=True, data_only=True)
    price_rows = list(workbook4.worksheets[0].iter_rows(min_row=2, values_only=True))
    workbook4.close()
    price_dates = np.asarray([_parse_date(row[0]) for row in price_rows], dtype="datetime64[D]")
    dynamic_price = np.asarray([row[1:] for row in price_rows], dtype=float)

    expected_dates = np.arange(np.datetime64("2025-01-01"), np.datetime64("2026-01-01"))
    if load_kw.shape != (365, 144) or actual_pv_kw.shape != (365, 144):
        raise ValueError("附件2负荷与光伏必须均为365×144")
    if dynamic_price.shape != (365, 144):
        raise ValueError(f"附件4动态电价必须为365×144，实际为{dynamic_price.shape}")
    if not np.array_equal(dates, expected_dates) or not np.array_equal(pv_dates, expected_dates):
        raise ValueError("附件2日期不连续或负荷、光伏日期不一致")
    if not np.array_equal(price_dates, expected_dates):
        raise ValueError("附件4日期与附件2不一致")
    if not np.array_equal(forecast_date_matrix, np.repeat(expected_dates[:, None], 4, axis=1)):
        raise ValueError("附件3日期与附件2不一致")
    if not np.array_equal(issue_matrix, np.repeat(np.asarray(RELEASE_HOURS)[None, :], 365, axis=0)):
        raise ValueError("附件3发布时刻必须依次为0、6、12、18点")
    arrays = {
        "负荷": load_kw,
        "实际光伏": actual_pv_kw,
        "光伏预测": forecast_array,
        "动态电价": dynamic_price,
    }
    for name, values in arrays.items():
        if not np.all(np.isfinite(values)) or np.any(values < 0.0):
            raise ValueError(f"{name}包含NaN/Inf或负值")
    return Problem4Data(
        dates,
        load_kw,
        actual_pv_kw,
        forecast_array,
        dynamic_price,
        tuple(path.name for path in (actual_path, forecast_path, price_path)),
    )


def autocorrelation(values: np.ndarray | Sequence[float], max_lag: int) -> np.ndarray:
    series = np.asarray(values, dtype=float).reshape(-1)
    if len(series) < 2 or max_lag < 0:
        raise ValueError("ACF序列过短或最大滞后非法")
    max_lag = min(int(max_lag), len(series) - 1)
    centered = series - series.mean()
    denominator = float(np.dot(centered, centered))
    if denominator <= 1e-20:
        return np.r_[1.0, np.zeros(max_lag)]
    return np.asarray(
        [1.0] + [float(np.dot(centered[:-lag], centered[lag:]) / denominator) for lag in range(1, max_lag + 1)]
    )


def analyze_pv_residuals(
    actual_forecast_pairs: Mapping[str, tuple[np.ndarray, np.ndarray]],
    max_acf_lag: int = 144,
) -> PVResidualAnalysis:
    """按e=实际光伏-预测光伏统计误差及ACF。"""
    statistics: dict[str, dict[str, float | int]] = {}
    residuals: dict[str, np.ndarray] = {}
    acfs: dict[str, np.ndarray] = {}
    for release, (actual, forecast) in actual_forecast_pairs.items():
        observed = np.asarray(actual, dtype=float)
        predicted = np.asarray(forecast, dtype=float)
        if observed.shape != predicted.shape or observed.size == 0:
            raise ValueError(f"{release}实际与预测数组必须同形且非空")
        residual = (observed - predicted).reshape(-1)
        residuals[release] = residual
        acfs[release] = autocorrelation(residual, max_acf_lag)
        statistics[release] = {
            "sample_count_10min": int(residual.size),
            "mae_kw": float(np.mean(np.abs(residual))),
            "rmse_kw": float(np.sqrt(np.mean(residual**2))),
            "residual_std_kw": float(np.std(residual)),
            "residual_mean_kw": float(np.mean(residual)),
        }
    ordered = [key for key in ("0:00", "6:00", "12:00", "18:00") if key in statistics]
    decreasing = bool(
        len(ordered) >= 2
        and all(
            float(statistics[left]["rmse_kw"]) >= float(statistics[right]["rmse_kw"])
            for left, right in zip(ordered, ordered[1:])
        )
    )
    return PVResidualAnalysis(statistics, residuals, acfs, decreasing)


def build_pv_actual_forecast_pairs(
    data: Problem4Data,
    start_day: int = OUTPUT_START_DAY,
) -> dict[str, tuple[np.ndarray, np.ndarray]]:
    """构造各发布时间至当日24:00的对齐样本，不跨越已发布信息边界。"""
    pairs: dict[str, tuple[np.ndarray, np.ndarray]] = {}
    for release_index, release_hour in enumerate(RELEASE_HOURS):
        actual_rows: list[np.ndarray] = []
        forecast_rows: list[np.ndarray] = []
        for day in range(start_day, len(data.dates)):
            anchor = (
                float(data.actual_pv_kw[day - 1, -1])
                if release_hour == 0 and day > 0
                else (0.0 if release_hour == 0 else float(data.actual_pv_kw[day, release_hour * 6 - 1]))
            )
            predicted = interpolate_pv_release(
                data.pv_forecast_hourly_kw[day, release_index], release_hour, anchor
            )
            actual = data.actual_pv_kw[day, release_hour * 6 :]
            if predicted.shape != actual.shape:
                raise RuntimeError("光伏实际与预测时间对齐失败")
            actual_rows.append(actual)
            forecast_rows.append(predicted)
        pairs[f"{release_hour}:00"] = (np.asarray(actual_rows), np.asarray(forecast_rows))
    return pairs


def problem2_strategy_metrics(
    result: Problem2YearResult,
    dynamic_price: np.ndarray,
    start_day: int = OUTPUT_START_DAY,
) -> dict[str, float | int | bool]:
    """以附件4逐日逐时电价独立重算问题2型策略的计划费和紧急购电费。"""
    price = np.asarray(dynamic_price, dtype=float)
    if price.shape != result.grid.shape:
        raise ValueError("动态电价与问题2结果必须同为days×144")
    selected = slice(start_day, len(result.grid))
    plan = result.grid[selected]
    emergency = result.emergency[selected]
    selected_price = price[selected]
    plan_cost = float(np.sum(selected_price * plan))
    emergency_cost = float(np.sum(5.0 * selected_price * emergency))
    actual_pv_total = result.pv_used[selected] + result.pv_curtailed[selected]
    return {
        "plan_purchase_kwh": float(plan.sum()),
        "actual_contract_grid_used_kwh": float(result.grid_used[selected].sum()),
        "emergency_purchase_kwh": float(emergency.sum()),
        "total_physical_grid_purchase_kwh": float(result.grid_used[selected].sum() + emergency.sum()),
        "plan_cost_yuan": plan_cost,
        "adjustment_cost_yuan": 0.0,
        "emergency_cost_yuan": emergency_cost,
        "total_cost_yuan": plan_cost + emergency_cost,
        "emergency_interval_count": int(np.sum(emergency > 1e-6)),
        "emergency_day_count": int(np.sum(emergency.sum(axis=1) > 1e-6)),
        "pv_consumed_kwh": float(result.pv_used[selected].sum()),
        "pv_curtailed_kwh": float(result.pv_curtailed[selected].sum()),
        "pv_available_kwh": float(actual_pv_total.sum()),
        "charge_kwh": float(result.charge[selected].sum()),
        "discharge_kwh": float(result.discharge[selected].sum()),
        "unused_contracted_energy_kwh": float(result.spill[selected].sum()),
    }


def storage_price_exposure(
    charge: np.ndarray,
    discharge: np.ndarray,
    dynamic_price: np.ndarray,
    start_day: int = OUTPUT_START_DAY,
) -> dict[str, float]:
    """用电量加权价格验证低价充电、高价放电，而非仅作定性描述。"""
    price = np.asarray(dynamic_price, dtype=float)[start_day:]
    charged = np.asarray(charge, dtype=float)[start_day:]
    discharged = np.asarray(discharge, dtype=float)[start_day:]
    charge_total = float(charged.sum())
    discharge_total = float(discharged.sum())
    average_charge_price = float(np.sum(price * charged) / max(charge_total, 1e-12))
    average_discharge_price = float(np.sum(price * discharged) / max(discharge_total, 1e-12))
    return {
        "charge_weighted_price_yuan_per_kwh": average_charge_price,
        "discharge_weighted_price_yuan_per_kwh": average_discharge_price,
        "discharge_minus_charge_price_yuan_per_kwh": average_discharge_price - average_charge_price,
        "charge_kwh": charge_total,
        "discharge_kwh": discharge_total,
    }


def _run_problem2_worker(arguments) -> Problem2ModelStudy:
    data, config = arguments
    return run_model_study(data, config=config, progress=False)


def _run_problem3_worker(arguments) -> StrategyResult:
    data, forecast, config = arguments
    return run_strategy(
        "C_动态电价+滚动MPC",
        data,
        forecast,
        (6, 12, 18),
        event_triggered=False,
        config=config,
        progress=False,
    )


def run_problem4_study(
    data: Problem4Data,
    config: DispatchConfig = DispatchConfig(),
    parallel: bool = True,
    progress: bool = True,
) -> Problem4Study:
    """复用问题2/3引擎完成A、B、C三策略全年回测。"""
    flat_value = float(np.mean(data.dynamic_price[OUTPUT_START_DAY:]))
    flat_price = np.full(144, flat_value)
    price_blind_data = data.as_problem2_data(flat_price)
    dynamic_problem2_data = data.as_problem2_data(data.dynamic_price)
    load_forecast = build_causal_load_forecasts(data.load_kw)
    dynamic_problem3_data = data.as_problem3_data(data.dynamic_price)
    residual_analysis = analyze_pv_residuals(build_pv_actual_forecast_pairs(data))

    if parallel:
        from concurrent.futures import ProcessPoolExecutor, as_completed

        jobs = {
            "A": (_run_problem2_worker, (price_blind_data, config)),
            "B": (_run_problem2_worker, (dynamic_problem2_data, config)),
            "C": (_run_problem3_worker, (dynamic_problem3_data, load_forecast, config)),
        }
        completed: dict[str, object] = {}
        with ProcessPoolExecutor(max_workers=3) as executor:
            futures = {
                executor.submit(function, arguments): name
                for name, (function, arguments) in jobs.items()
            }
            for future in as_completed(futures):
                name = futures[future]
                completed[name] = future.result()
                if progress:
                    print(f"strategy {name}: annual simulation complete", flush=True)
        price_blind_study = completed["A"]
        dynamic_problem2_study = completed["B"]
        dynamic_mpc_result = completed["C"]
    else:
        price_blind_study = run_model_study(price_blind_data, config, progress)
        dynamic_problem2_study = run_model_study(dynamic_problem2_data, config, progress)
        dynamic_mpc_result = _run_problem3_worker((dynamic_problem3_data, load_forecast, config))
    if not isinstance(price_blind_study, Problem2ModelStudy):
        raise TypeError("策略A返回类型错误")
    if not isinstance(dynamic_problem2_study, Problem2ModelStudy):
        raise TypeError("策略B返回类型错误")
    if not isinstance(dynamic_mpc_result, StrategyResult):
        raise TypeError("策略C返回类型错误")

    comparison = {
        "A_不考虑动态电价": problem2_strategy_metrics(
            price_blind_study.final_result, data.dynamic_price
        ),
        "B_动态电价+储能优化": problem2_strategy_metrics(
            dynamic_problem2_study.final_result, data.dynamic_price
        ),
        "C_动态电价+滚动MPC": strategy_metrics(dynamic_mpc_result),
    }
    baseline_cost = float(comparison["A_不考虑动态电价"]["total_cost_yuan"])
    for values in comparison.values():
        values["saving_vs_original_ratio"] = float(
            (baseline_cost - float(values["total_cost_yuan"])) / baseline_cost
        )
    arbitrage = {
        "A_不考虑动态电价": storage_price_exposure(
            price_blind_study.final_result.charge,
            price_blind_study.final_result.discharge,
            data.dynamic_price,
        ),
        "B_动态电价+储能优化": storage_price_exposure(
            dynamic_problem2_study.final_result.charge,
            dynamic_problem2_study.final_result.discharge,
            data.dynamic_price,
        ),
        "C_动态电价+滚动MPC": storage_price_exposure(
            dynamic_mpc_result.charge,
            dynamic_mpc_result.discharge,
            data.dynamic_price,
        ),
    }

    audit42 = audit_problem2(dynamic_problem2_study.final_result, dynamic_problem2_data, config)
    audit42.pop("attachment1_price_repeated_daily", None)
    audit42.pop("attachments3_and_4_used", None)
    audit42.update(
        {
            "input_whitelist": list(data.source_files),
            "price_source": "附件4.xlsx",
            "attachment1_opened": False,
            "attachment2_opened": True,
            "attachment3_used_in_problem4_2_optimization": False,
            "attachment4_opened": True,
            "dynamic_price_shape_ok": bool(data.dynamic_price.shape == (365, 144)),
            "future_data_leakage": bool(dynamic_problem2_study.forecast_suite.leakage_detected),
            "throughput_penalty_lambda": config.throughput_penalty,
        }
    )
    audit43 = audit_problem3(dynamic_mpc_result, dynamic_problem3_data, load_forecast, config)
    audit43.update(
        {
            "input_whitelist": list(data.source_files),
            "price_source": "附件4.xlsx",
            "actual_source": "附件2.xlsx",
            "pv_forecast_source": "附件3.xlsx",
            "attachment4_opened": True,
            "dynamic_price_shape_ok": bool(data.dynamic_price.shape == (365, 144)),
            "future_data_leakage": bool(
                load_forecast.leakage_detected
                or bool(audit43["pv_future_actual_used_in_optimization"])
            ),
            "throughput_penalty_lambda": config.throughput_penalty,
        }
    )
    return Problem4Study(
        price_blind_study,
        dynamic_problem2_study,
        dynamic_mpc_result,
        load_forecast,
        residual_analysis,
        comparison,
        arbitrage,
        audit42,
        audit43,
    )


def _rolling_result(result: StrategyResult) -> RollingYearResult:
    return RollingYearResult(
        result.plan_grid,
        result.adjusted_grid,
        result.charge,
        result.discharge,
        result.pv_used,
        result.emergency,
        result.soc,
        result.daily_plan_cost,
        result.daily_adjustment_cost,
        result.daily_emergency_cost,
        result.daily_downward_kwh,
        result.daily_upward_kwh,
    )


def write_problem4_workbooks(
    root: str | Path,
    study: Problem4Study,
    output42: str | Path,
    output43: str | Path,
    dynamic_price: np.ndarray,
) -> tuple[Path, Path]:
    path42 = save_result2(
        root,
        output42,
        study.dynamic_problem2_study.final_result,
        template_name="result4-2.xlsx",
    )
    path43 = write_rolling_workbook(
        root,
        output43,
        "result4-3.xlsx",
        _rolling_result(study.dynamic_mpc_result),
        dynamic_price,
    )
    return path42, path43


def create_problem4_figures(
    root: str | Path,
    data: Problem4Data,
    study: Problem4Study,
) -> list[Path]:
    output_dir = Path(root).resolve() / "problem4_figures"
    output_dir.mkdir(parents=True, exist_ok=True)
    pairs = build_pv_actual_forecast_pairs(data)
    try:
        import matplotlib
    except ModuleNotFoundError:
        from problem4_plotting import create_problem4_svg_figures

        return create_problem4_svg_figures(
            output_dir, data, study, pairs, OUTPUT_START_DAY
        )

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    colors = {"0:00": "#355C7D", "6:00": "#2A9D8F", "12:00": "#E9C46A", "18:00": "#E76F51"}

    fig, axes = plt.subplots(4, 1, figsize=(12, 10), sharex=True)
    dates = data.dates[OUTPUT_START_DAY:].astype("datetime64[D]").astype(object)
    for axis, release in zip(axes, ("0:00", "6:00", "12:00", "18:00")):
        actual, forecast = pairs[release]
        daily_mean = np.mean(actual - forecast, axis=1)
        axis.plot(dates, daily_mean, color=colors[release], linewidth=0.8)
        axis.axhline(0.0, color="#555555", linewidth=0.7)
        axis.set_ylabel(f"{release}\nresidual (kW)")
        axis.grid(alpha=0.18)
    axes[-1].set_xlabel("Date")
    fig.suptitle("Daily mean PV forecast residual: actual minus forecast")
    fig.tight_layout(rect=(0, 0, 1, 0.97))
    residual_path = output_dir / "pv_residual_timeseries.png"
    fig.savefig(residual_path, dpi=180, bbox_inches="tight")
    plt.close(fig)

    fig, axes = plt.subplots(2, 2, figsize=(12, 7), sharex=True, sharey=True)
    for axis, release in zip(axes.ravel(), ("0:00", "6:00", "12:00", "18:00")):
        acf = study.residual_analysis.acf[release]
        axis.vlines(np.arange(len(acf)), 0.0, acf, color=colors[release], linewidth=0.8)
        axis.axhline(0.0, color="#555555", linewidth=0.7)
        axis.set_title(f"Issue {release}")
        axis.set_xlabel("Lag (10-min intervals)")
        axis.set_ylabel("ACF")
        axis.grid(alpha=0.18)
    fig.suptitle("Autocorrelation of PV forecast residuals")
    fig.tight_layout(rect=(0, 0, 1, 0.95))
    acf_path = output_dir / "pv_residual_acf.png"
    fig.savefig(acf_path, dpi=180, bbox_inches="tight")
    plt.close(fig)

    names = list(study.strategy_comparison)
    short_names = ["A price-blind", "B dynamic", "C dynamic MPC"]
    plan_cost = np.asarray([float(study.strategy_comparison[name]["plan_cost_yuan"]) for name in names]) / 1e6
    adjustment_cost = np.asarray([float(study.strategy_comparison[name].get("adjustment_cost_yuan", 0.0)) for name in names]) / 1e6
    emergency_cost = np.asarray([float(study.strategy_comparison[name]["emergency_cost_yuan"]) for name in names]) / 1e6
    fig, axis = plt.subplots(figsize=(9, 5))
    x = np.arange(len(names))
    axis.bar(x, plan_cost, color="#355C7D", label="Plan")
    axis.bar(x, adjustment_cost, bottom=plan_cost, color="#E9C46A", label="Adjustment")
    axis.bar(x, emergency_cost, bottom=plan_cost + adjustment_cost, color="#E76F51", label="Emergency")
    axis.set_xticks(x, short_names)
    axis.set_ylabel("Cost (million yuan)")
    axis.set_title("Problem 4 strategy cost comparison")
    axis.legend(frameon=False, ncol=3, loc="upper center")
    axis.grid(axis="y", alpha=0.18)
    fig.tight_layout()
    comparison_path = output_dir / "strategy_cost_comparison.png"
    fig.savefig(comparison_path, dpi=180, bbox_inches="tight")
    plt.close(fig)

    output_prices = data.dynamic_price[OUTPUT_START_DAY:]
    representative_day = OUTPUT_START_DAY + int(np.argmax(np.ptp(output_prices, axis=1)))
    interval = np.arange(144) / 6.0
    b_result = study.dynamic_problem2_study.final_result
    c_result = study.dynamic_mpc_result
    fig, axes = plt.subplots(3, 1, figsize=(12, 9), sharex=True)
    axes[0].plot(interval, data.dynamic_price[representative_day], color="#333333", linewidth=1.2)
    axes[0].set_ylabel("Price\n(yuan/kWh)")
    axes[0].set_title(f"Dynamic price and storage dispatch on {str(data.dates[representative_day])}")
    actions = (
        (b_result.charge[representative_day] - b_result.discharge[representative_day], "B dynamic", "#2A9D8F"),
        (c_result.charge[representative_day] - c_result.discharge[representative_day], "C dynamic MPC", "#E76F51"),
    )
    for action, label, color in actions:
        axes[1].plot(interval, action * 6.0, label=label, color=color, linewidth=1.0)
    axes[1].axhline(0.0, color="#555555", linewidth=0.7)
    axes[1].set_ylabel("Battery power\n(kW; + charge)")
    axes[1].legend(frameon=False, ncol=2)
    axes[2].plot(np.arange(145) / 6.0, b_result.soc[representative_day], label="B dynamic", color="#2A9D8F")
    axes[2].plot(np.arange(145) / 6.0, c_result.soc[representative_day], label="C dynamic MPC", color="#E76F51")
    axes[2].axhline(1200.0, color="#777777", linestyle="--", linewidth=0.7)
    axes[2].axhline(10800.0, color="#777777", linestyle="--", linewidth=0.7)
    axes[2].set_ylabel("SOC (kWh)")
    axes[2].set_xlabel("Hour")
    axes[2].legend(frameon=False, ncol=2)
    for axis in axes:
        axis.grid(alpha=0.18)
    fig.tight_layout()
    soc_path = output_dir / "dynamic_price_storage_soc.png"
    fig.savefig(soc_path, dpi=180, bbox_inches="tight")
    plt.close(fig)
    return [residual_path, acf_path, comparison_path, soc_path]


def _jsonable(value):
    if isinstance(value, dict):
        return {str(key): _jsonable(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_jsonable(item) for item in value]
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, (np.integer, np.floating)):
        return value.item()
    if isinstance(value, np.bool_):
        return bool(value)
    return value


def write_problem4_reports(
    root: str | Path,
    data: Problem4Data,
    study: Problem4Study,
    figures: Sequence[Path],
) -> tuple[Path, Path, Path, Path]:
    import csv
    import json

    root_path = Path(root).resolve()
    json_path = root_path / "problem4_analysis.json"
    report_path = root_path / "problem4_analysis_report.md"
    comparison_csv = root_path / "problem4_strategy_comparison.csv"
    residual_csv = root_path / "problem4_pv_residual_statistics.csv"
    payload = {
        "input_whitelist": list(data.source_files),
        "output_period": "2025-02-01 to 2025-12-31",
        "time_alignment": {
            "attachment2_and_4_columns_are_interval_endpoints": True,
            "first_interval": "0:00-0:10",
            "last_interval": "23:50-0:00+1",
            "attachment3_prediction_1h": "first clock-hour endpoint after issue time",
            "pv_interpolation": "linear interpolation from the observed issue-time anchor to future hourly endpoints",
        },
        "physical_constraints": {
            "balance": "grid_used + emergency + pv - curtailment + discharge = load + charge",
            "soc_bounds_kwh": [1200.0, 10800.0],
            "power_limit_kw": 5000.0,
            "interval_limit_kwh": 5000.0 / 6.0,
            "efficiency": 0.9,
            "no_reverse_sale": True,
            "no_simultaneous_charge_discharge": True,
            "throughput_penalty_lambda": DispatchConfig().throughput_penalty,
        },
        "problem4_2_selected_forecast_model": study.dynamic_problem2_study.selected_model,
        "price_blind_selected_forecast_model": study.price_blind_study.selected_model,
        "pv_residual_statistics": study.residual_analysis.statistics,
        "error_decreases_with_shorter_horizon": study.residual_analysis.error_decreases_with_shorter_horizon,
        "strategy_comparison": study.strategy_comparison,
        "arbitrage_analysis": study.arbitrage_analysis,
        "audit_problem4_2": study.audit_problem4_2,
        "audit_problem4_3": study.audit_problem4_3,
        "figures": [path.name for path in figures],
        "cvar_decision": "not retained: the deterministic causal MPC is stable and no validated scenario set demonstrated lower tail cost after adjustment fees",
    }
    json_path.write_text(json.dumps(_jsonable(payload), ensure_ascii=False, indent=2), encoding="utf-8")

    all_fields = list(dict.fromkeys(
        ["strategy"] + [key for values in study.strategy_comparison.values() for key in values]
    ))
    with comparison_csv.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=all_fields)
        writer.writeheader()
        for name, values in study.strategy_comparison.items():
            writer.writerow({"strategy": name, **values})
    with residual_csv.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["release", "sample_count_10min", "mae_kw", "rmse_kw", "residual_std_kw", "residual_mean_kw"],
        )
        writer.writeheader()
        for release, values in study.residual_analysis.statistics.items():
            writer.writerow({"release": release, **values})

    lines = [
        "# 问题四：动态电价感知与滚动MPC优化",
        "",
        "数据仅来自附件2、附件3、附件4。问题4-2复用问题2的因果预测、分位数校准和实时回放；"
        "问题4-3复用问题3的0/6/12/18点滚动MPC和非对称调整费用。",
        "",
        "## 光伏预测残差",
        "",
        "残差定义为 e=实际光伏-预测光伏。统计范围为各发布时间至当日24:00的可执行时域。",
        "",
        "| 发布时间 | 样本数 | MAE/kW | RMSE/kW | 残差标准差/kW | 残差均值/kW |",
        "|---|---:|---:|---:|---:|---:|",
    ]
    for release, values in study.residual_analysis.statistics.items():
        lines.append(
            f"| {release} | {values['sample_count_10min']} | {values['mae_kw']:.3f} | "
            f"{values['rmse_kw']:.3f} | {values['residual_std_kw']:.3f} | {values['residual_mean_kw']:.3f} |"
        )
    lines.extend(
        [
            "",
            "RMSE随发布时间由0:00推进到18:00单调下降。预测时距缩短确实提高了精度，支持采用多阶段滚动修正。"
            if study.residual_analysis.error_decreases_with_shorter_horizon
            else "RMSE未呈严格单调下降，应谨慎评价新增发布时间的价值。",
            "",
            "## 策略比较（2025-02-01至2025-12-31）",
            "",
            "| 策略 | 总费用/元 | 计划费用/元 | 调整费用/元 | 紧急购电费/元 | 紧急电量/kWh | 弃光/kWh | 充电/kWh | 放电/kWh | 相对A节省 |",
            "|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|",
        ]
    )
    for name, values in study.strategy_comparison.items():
        lines.append(
            f"| {name} | {values['total_cost_yuan']:.2f} | {values['plan_cost_yuan']:.2f} | "
            f"{values.get('adjustment_cost_yuan', 0.0):.2f} | {values['emergency_cost_yuan']:.2f} | "
            f"{values['emergency_purchase_kwh']:.2f} | {values['pv_curtailed_kwh']:.2f} | "
            f"{values['charge_kwh']:.2f} | {values['discharge_kwh']:.2f} | "
            f"{100.0 * values['saving_vs_original_ratio']:.3f}% |"
        )
    lines.extend(
        [
            "",
            "## 峰谷套利验证",
            "",
            "| 策略 | 充电加权电价/元·kWh⁻¹ | 放电加权电价/元·kWh⁻¹ | 放电-充电价差/元·kWh⁻¹ |",
            "|---|---:|---:|---:|",
        ]
    )
    for name, values in study.arbitrage_analysis.items():
        lines.append(
            f"| {name} | {values['charge_weighted_price_yuan_per_kwh']:.4f} | "
            f"{values['discharge_weighted_price_yuan_per_kwh']:.4f} | "
            f"{values['discharge_minus_charge_price_yuan_per_kwh']:.4f} |"
        )
    lines.extend(
        [
            "",
            "## 数值审计",
            "",
            f"- 问题4-2全部检查通过：{study.audit_problem4_2['all_checks_passed']}。",
            f"- 问题4-3全部检查通过：{study.audit_problem4_3['all_required_checks_passed']}。",
            f"- 问题4-2最大平衡误差：{study.audit_problem4_2['max_power_balance_error_kwh']:.3e} kWh。",
            f"- 问题4-3最大平衡误差：{study.audit_problem4_3['max_power_balance_error_kwh']:.3e} kWh。",
            f"- 问题4-3当前SOC继承误差：{study.audit_problem4_3['max_update_soc_inheritance_error_kwh']:.3e} kWh。",
            "- CVaR未强行保留：在没有验证场景显示尾部成本改善前，避免增加不必要的风险参数。",
        ]
    )
    report_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return json_path, report_path, comparison_csv, residual_csv


def print_key_results(study: Problem4Study) -> None:
    import json

    print("\nProblem 4 strategy comparison (2025-02-01 to 2025-12-31)")
    for name, values in study.strategy_comparison.items():
        print(
            f"{name}: total={values['total_cost_yuan']:.2f}, plan={values['plan_cost_yuan']:.2f}, "
            f"adjust={values.get('adjustment_cost_yuan', 0.0):.2f}, "
            f"emergency={values['emergency_cost_yuan']:.2f}, "
            f"saving={100.0 * values['saving_vs_original_ratio']:.3f}%"
        )
    print("\nPV residual statistics")
    print(json.dumps(_jsonable(study.residual_analysis.statistics), ensure_ascii=False, indent=2))
    print("\nProblem 4-2 audit")
    print(json.dumps(_jsonable(study.audit_problem4_2), ensure_ascii=False, indent=2))
    print("\nProblem 4-3 audit")
    print(json.dumps(_jsonable(study.audit_problem4_3), ensure_ascii=False, indent=2))


def main() -> None:
    import argparse

    parser = argparse.ArgumentParser(description="问题四：动态电价 + 储能套利 + 滚动MPC")
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parent)
    parser.add_argument("--output42", type=Path, default=None)
    parser.add_argument("--output43", type=Path, default=None)
    parser.add_argument("--no-parallel", action="store_true")
    parser.add_argument("--quiet", action="store_true")
    args = parser.parse_args()
    root = args.root.resolve()
    output42 = (args.output42 or root / "result4-2.xlsx").resolve()
    output43 = (args.output43 or root / "result4-3.xlsx").resolve()
    data = load_problem4_data(root, print_structure=not args.quiet)
    study = run_problem4_study(
        data,
        parallel=not args.no_parallel,
        progress=not args.quiet,
    )
    write_problem4_workbooks(root, study, output42, output43, data.dynamic_price)
    figures = create_problem4_figures(root, data, study)
    write_problem4_reports(root, data, study, figures)
    print_key_results(study)


if __name__ == "__main__":
    main()

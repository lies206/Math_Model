"""2026 数学建模 C 题：微网与外部电网电力调控策略。

从当前目录的附件读取数据，使用 SciPy HiGHS 构造稀疏线性规划，
并在不改变模板工作表名、表头和已有样式的前提下生成五个结果文件。
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Sequence

import numpy as np


@dataclass(frozen=True)
class DispatchConfig:
    interval_hours: float = 1.0 / 6.0
    efficiency: float = 0.90
    soc_min_kwh: float = 1200.0
    soc_max_kwh: float = 10800.0
    power_limit_kw: float = 5000.0
    throughput_penalty: float = 1e-8

    @property
    def interval_limit_kwh(self) -> float:
        return self.power_limit_kw * self.interval_hours


@dataclass(frozen=True)
class AdjustmentSettlement:
    downward_kwh: float
    upward_kwh: float
    penalty_cost: float
    incremental_cost: float


@dataclass
class DispatchResult:
    success: bool
    message: str
    grid: np.ndarray
    charge: np.ndarray
    discharge: np.ndarray
    pv_used: np.ndarray
    soc: np.ndarray
    purchase_cost: float


@dataclass(frozen=True)
class InputData:
    static_price: np.ndarray
    problem1_load_kw: np.ndarray
    problem1_pv_kw: np.ndarray
    dates: np.ndarray
    load_kw: np.ndarray
    actual_pv_kw: np.ndarray
    forecast_pv_kw: np.ndarray
    dynamic_price: np.ndarray


@dataclass
class RollingDayResult:
    plan: DispatchResult
    final_grid: np.ndarray
    charge: np.ndarray
    discharge: np.ndarray
    pv_used_actual: np.ndarray
    emergency: np.ndarray
    soc: np.ndarray
    adjustment: AdjustmentSettlement
    emergency_cost: float
    total_cost: float


@dataclass
class DeterministicYearResult:
    grid: np.ndarray
    charge: np.ndarray
    discharge: np.ndarray
    pv_used: np.ndarray
    emergency: np.ndarray
    soc: np.ndarray
    daily_purchase_cost: np.ndarray
    total_cost: float


@dataclass
class RollingYearResult:
    plan_grid: np.ndarray
    adjusted_grid: np.ndarray
    charge: np.ndarray
    discharge: np.ndarray
    pv_used_actual: np.ndarray
    emergency: np.ndarray
    soc: np.ndarray
    daily_plan_cost: np.ndarray
    daily_adjustment_cost: np.ndarray
    daily_emergency_cost: np.ndarray
    downward_kwh: np.ndarray
    upward_kwh: np.ndarray

    @property
    def daily_total_cost(self) -> np.ndarray:
        return self.daily_plan_cost + self.daily_adjustment_cost + self.daily_emergency_cost

    @property
    def total_cost(self) -> float:
        return float(self.daily_total_cost.sum())


def kw_to_interval_kwh(values: np.ndarray | Sequence[float], interval_hours: float = 1.0 / 6.0) -> np.ndarray:
    """把区间平均功率 kW 转换为区间电量 kWh。"""
    return np.asarray(values, dtype=float) * interval_hours


def expand_hourly_forecast(values: np.ndarray | Sequence[float]) -> np.ndarray:
    """把未来整小时预测按零阶保持展开；第1个值覆盖发布时刻后的第1小时。"""
    return np.repeat(np.asarray(values, dtype=float), 6)


def _sheet_matrix(sheet, min_row: int, max_row: int, min_col: int, max_col: int) -> np.ndarray:
    rows = sheet.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
        values_only=True,
    )
    return np.asarray(list(rows), dtype=float)


def load_input_data(root: str | Path) -> InputData:
    """读取附件 1-4，并校验题目规定的日期和矩阵形状。"""
    from datetime import datetime
    from openpyxl import load_workbook

    root_path = Path(root).resolve()
    attachment_dir = root_path / "附件"

    wb1 = load_workbook(attachment_dir / "附件1.xlsx", read_only=True, data_only=True)
    ws1 = wb1[wb1.sheetnames[0]]
    static_price = _sheet_matrix(ws1, 2, 145, 2, 2).reshape(-1)
    problem1_load = _sheet_matrix(ws1, 2, 145, 3, 3).reshape(-1)
    problem1_pv = _sheet_matrix(ws1, 2, 145, 4, 4).reshape(-1)
    wb1.close()

    wb2 = load_workbook(attachment_dir / "附件2.xlsx", read_only=True, data_only=True)
    load_sheet = wb2["小区负载"]
    pv_sheet = wb2["光伏发电实际功率"]
    date_values = [row[0] for row in load_sheet.iter_rows(min_row=2, max_row=366, min_col=1, max_col=1, values_only=True)]
    dates = np.asarray([np.datetime64(value.date()) for value in date_values], dtype="datetime64[D]")
    load_kw = _sheet_matrix(load_sheet, 2, 366, 2, 145)
    actual_pv_kw = _sheet_matrix(pv_sheet, 2, 366, 2, 145)
    wb2.close()

    wb3 = load_workbook(attachment_dir / "附件3.xlsx", read_only=True, data_only=True)
    forecast_sheet = wb3[wb3.sheetnames[0]]
    forecast = np.empty((365, 4, 24), dtype=float)
    expected_releases = ("0:00", "6:00", "12:00", "18:00")
    forecast_rows = forecast_sheet.iter_rows(min_row=2, max_row=1461, min_col=1, max_col=26, values_only=True)
    for row_index, row_values in enumerate(forecast_rows):
        day, release_index = divmod(row_index, 4)
        excel_row = row_index + 2
        if release_index == 0:
            parsed = datetime.strptime(str(row_values[0]), "%Y-%m-%d")
            if np.datetime64(parsed.date()) != dates[day]:
                raise ValueError(f"附件3第 {excel_row} 行日期与附件2不一致")
        expected = expected_releases[release_index]
        if str(row_values[1]) != expected:
            raise ValueError(f"附件3第 {excel_row} 行预报时刻应为 {expected}")
        forecast[day, release_index, :] = row_values[2:26]
    wb3.close()

    wb4 = load_workbook(attachment_dir / "附件4.xlsx", read_only=True, data_only=True)
    price_sheet = wb4[wb4.sheetnames[0]]
    dynamic_price = _sheet_matrix(price_sheet, 2, 366, 2, 145)
    wb4.close()

    expected_dates = np.arange(np.datetime64("2025-01-01"), np.datetime64("2026-01-01"))
    if not np.array_equal(dates, expected_dates):
        raise ValueError("附件2日期必须完整覆盖 2025-01-01 至 2025-12-31")
    arrays = (static_price, problem1_load, problem1_pv, load_kw, actual_pv_kw, forecast, dynamic_price)
    if not all(np.all(np.isfinite(array)) and np.all(array >= 0.0) for array in arrays):
        raise ValueError("附件数据包含缺失、非数值或负值")

    return InputData(
        static_price=static_price,
        problem1_load_kw=problem1_load,
        problem1_pv_kw=problem1_pv,
        dates=dates,
        load_kw=load_kw,
        actual_pv_kw=actual_pv_kw,
        forecast_pv_kw=forecast,
        dynamic_price=dynamic_price,
    )


def soc_path(initial_soc: float, charge: np.ndarray, discharge: np.ndarray, efficiency: float) -> np.ndarray:
    """按交流母线侧充放电量计算含初值的完整 SOC 路径。"""
    charge = np.asarray(charge, dtype=float)
    discharge = np.asarray(discharge, dtype=float)
    if charge.shape != discharge.shape:
        raise ValueError("charge and discharge must have the same shape")
    increments = efficiency * charge - discharge / efficiency
    return np.r_[float(initial_soc), float(initial_soc) + np.cumsum(increments)]


def adjustment_settlement(plan: np.ndarray, adjusted: np.ndarray, price: np.ndarray) -> AdjustmentSettlement:
    """计算下调违约量、上调超购量及其题面规定的附加费用。"""
    plan = np.asarray(plan, dtype=float)
    adjusted = np.asarray(adjusted, dtype=float)
    price = np.asarray(price, dtype=float)
    downward = np.maximum(plan - adjusted, 0.0)
    upward = np.maximum(adjusted - plan, 0.0)
    return AdjustmentSettlement(
        downward_kwh=float(downward.sum()),
        upward_kwh=float(upward.sum()),
        penalty_cost=float(np.dot(0.5 * price, downward)),
        incremental_cost=float(np.dot(1.5 * price, upward)),
    )


def _clock_label(total_minutes: int, mark_next_day: bool) -> str:
    day = total_minutes // 1440
    minute_of_day = total_minutes % 1440
    hour, minute = divmod(minute_of_day, 60)
    suffix = "+1" if mark_next_day and day >= 1 else ""
    return f"{hour}:{minute:02d}{suffix}"


def build_time_labels() -> list[str]:
    """返回当天 0:00 至 24:00 的 144 个左闭右开 10 分钟区间。"""
    labels: list[str] = []
    for index in range(144):
        start = 10 * index
        end = start + 10
        labels.append(f"{_clock_label(start, False)}-{_clock_label(end, True)}")
    return labels


def group_emergency_intervals(amounts: np.ndarray | Sequence[float], start_minutes: int = 0, tolerance: float = 1e-7) -> list[tuple[str, float]]:
    """把相邻的非零紧急购电区间合并为模板要求的时间段。"""
    values = np.asarray(amounts, dtype=float)
    groups: list[tuple[str, float]] = []
    index = 0
    while index < len(values):
        if values[index] <= tolerance:
            index += 1
            continue
        end_index = index + 1
        while end_index < len(values) and values[end_index] > tolerance:
            end_index += 1
        start = start_minutes + 10 * index
        end = start_minutes + 10 * end_index
        label = f"{_clock_label(start, False)}-{_clock_label(end, True)}"
        groups.append((label, float(values[index:end_index].sum())))
        index = end_index
    return groups


def solve_dispatch(
    load_kwh: np.ndarray | Sequence[float],
    pv_kwh: np.ndarray | Sequence[float],
    price: np.ndarray | Sequence[float],
    initial_soc: float,
    terminal_soc: float | None,
    config: DispatchConfig = DispatchConfig(),
) -> DispatchResult:
    """求解确定性购电、光伏消纳和储能调度稀疏线性规划。"""
    from scipy.optimize import linprog
    from scipy.sparse import coo_matrix

    load = np.asarray(load_kwh, dtype=float).reshape(-1)
    pv = np.asarray(pv_kwh, dtype=float).reshape(-1)
    prices = np.asarray(price, dtype=float).reshape(-1)
    if not (len(load) == len(pv) == len(prices)):
        raise ValueError("load, pv and price must have the same length")
    if len(load) == 0 or np.any(load < 0) or np.any(pv < 0) or np.any(prices < 0):
        raise ValueError("input arrays must be nonempty and nonnegative")
    if not (config.soc_min_kwh <= initial_soc <= config.soc_max_kwh):
        raise ValueError("initial SOC is outside the permitted range")
    if terminal_soc is not None and not (config.soc_min_kwh <= terminal_soc <= config.soc_max_kwh):
        raise ValueError("terminal SOC is outside the permitted range")

    periods = len(load)
    grid_offset = 0
    charge_offset = periods
    discharge_offset = 2 * periods
    pv_offset = 3 * periods
    soc_offset = 4 * periods
    variable_count = 5 * periods + 1

    objective = np.zeros(variable_count)
    objective[grid_offset : grid_offset + periods] = prices
    objective[charge_offset : discharge_offset + periods] = config.throughput_penalty

    rows: list[int] = []
    columns: list[int] = []
    values: list[float] = []
    for t in range(periods):
        rows.extend((t, t, t, t))
        columns.extend((grid_offset + t, charge_offset + t, discharge_offset + t, pv_offset + t))
        values.extend((1.0, -1.0, 1.0, 1.0))

        row = periods + t
        rows.extend((row, row, row, row))
        columns.extend((soc_offset + t, soc_offset + t + 1, charge_offset + t, discharge_offset + t))
        values.extend((-1.0, 1.0, -config.efficiency, 1.0 / config.efficiency))

    equality_matrix = coo_matrix(
        (values, (rows, columns)), shape=(2 * periods, variable_count)
    ).tocsr()
    equality_rhs = np.r_[load, np.zeros(periods)]

    lower = np.zeros(variable_count)
    upper = np.full(variable_count, np.inf)
    upper[charge_offset : discharge_offset] = config.interval_limit_kwh
    upper[discharge_offset : pv_offset] = config.interval_limit_kwh
    upper[pv_offset : soc_offset] = pv
    lower[soc_offset:] = config.soc_min_kwh
    upper[soc_offset:] = config.soc_max_kwh
    lower[soc_offset] = upper[soc_offset] = float(initial_soc)
    if terminal_soc is not None:
        lower[-1] = upper[-1] = float(terminal_soc)

    result = linprog(
        objective,
        A_eq=equality_matrix,
        b_eq=equality_rhs,
        bounds=np.column_stack((lower, upper)),
        method="highs",
        options={"presolve": True},
    )
    if not result.success:
        empty = np.full(periods, np.nan)
        return DispatchResult(False, result.message, empty, empty.copy(), empty.copy(), empty.copy(), np.full(periods + 1, np.nan), float("nan"))

    solution = result.x
    grid = solution[grid_offset : charge_offset]
    charge = solution[charge_offset : discharge_offset]
    discharge = solution[discharge_offset : pv_offset]
    pv_used = solution[pv_offset : soc_offset]
    soc = solution[soc_offset:]
    return DispatchResult(
        success=True,
        message=result.message,
        grid=grid,
        charge=charge,
        discharge=discharge,
        pv_used=pv_used,
        soc=soc,
        purchase_cost=float(np.dot(prices, grid)),
    )


def _solve_adjusted_dispatch(
    load_kwh: np.ndarray,
    pv_kwh: np.ndarray,
    price: np.ndarray,
    reference_grid: np.ndarray,
    initial_soc: float,
    terminal_soc: float,
    config: DispatchConfig,
) -> DispatchResult:
    """以原计划为基准，按下调 50%、上调 150% 的费用求解剩余时域。"""
    from scipy.optimize import Bounds, LinearConstraint, milp
    from scipy.sparse import coo_matrix

    load = np.asarray(load_kwh, dtype=float)
    pv = np.asarray(pv_kwh, dtype=float)
    prices = np.asarray(price, dtype=float)
    reference = np.asarray(reference_grid, dtype=float)
    periods = len(load)
    if not (pv.shape == prices.shape == reference.shape == load.shape):
        raise ValueError("rolling arrays must have identical one-dimensional shapes")

    grid_offset = 0
    charge_offset = periods
    discharge_offset = 2 * periods
    pv_offset = 3 * periods
    soc_offset = 4 * periods
    down_offset = 5 * periods + 1
    up_offset = 6 * periods + 1
    mode_offset = 7 * periods + 1
    variable_count = 8 * periods + 1

    objective = np.zeros(variable_count)
    objective[charge_offset : discharge_offset + periods] = config.throughput_penalty
    objective[down_offset : down_offset + periods] = 0.5 * prices
    objective[up_offset : up_offset + periods] = 1.5 * prices

    rows: list[int] = []
    columns: list[int] = []
    values: list[float] = []
    for t in range(periods):
        rows.extend((t, t, t, t))
        columns.extend((grid_offset + t, charge_offset + t, discharge_offset + t, pv_offset + t))
        values.extend((1.0, -1.0, 1.0, 1.0))

        soc_row = periods + t
        rows.extend((soc_row, soc_row, soc_row, soc_row))
        columns.extend((soc_offset + t, soc_offset + t + 1, charge_offset + t, discharge_offset + t))
        values.extend((-1.0, 1.0, -config.efficiency, 1.0 / config.efficiency))

        adjustment_row = 2 * periods + t
        rows.extend((adjustment_row, adjustment_row, adjustment_row))
        columns.extend((grid_offset + t, down_offset + t, up_offset + t))
        values.extend((1.0, 1.0, -1.0))

        charge_mode_row = 3 * periods + t
        rows.extend((charge_mode_row, charge_mode_row))
        columns.extend((charge_offset + t, mode_offset + t))
        values.extend((1.0, -config.interval_limit_kwh))

        discharge_mode_row = 4 * periods + t
        rows.extend((discharge_mode_row, discharge_mode_row))
        columns.extend((discharge_offset + t, mode_offset + t))
        values.extend((1.0, config.interval_limit_kwh))

    matrix = coo_matrix((values, (rows, columns)), shape=(5 * periods, variable_count)).tocsr()
    rhs = np.r_[load, np.zeros(periods), reference]
    constraint_lower = np.r_[rhs, np.full(2 * periods, -np.inf)]
    constraint_upper = np.r_[rhs, np.zeros(periods), np.full(periods, config.interval_limit_kwh)]

    lower = np.zeros(variable_count)
    upper = np.full(variable_count, np.inf)
    upper[charge_offset : discharge_offset] = config.interval_limit_kwh
    upper[discharge_offset : pv_offset] = config.interval_limit_kwh
    upper[pv_offset : soc_offset] = pv
    lower[soc_offset : down_offset] = config.soc_min_kwh
    upper[soc_offset : down_offset] = config.soc_max_kwh
    lower[soc_offset] = upper[soc_offset] = float(initial_soc)
    lower[down_offset - 1] = upper[down_offset - 1] = float(terminal_soc)
    upper[mode_offset:] = 1.0

    integrality = np.zeros(variable_count, dtype=int)
    integrality[mode_offset:] = 1

    result = milp(
        objective,
        integrality=integrality,
        bounds=Bounds(lower, upper),
        constraints=LinearConstraint(matrix, constraint_lower, constraint_upper),
        options={"presolve": True},
    )
    if not result.success:
        empty = np.full(periods, np.nan)
        return DispatchResult(False, result.message, empty, empty.copy(), empty.copy(), empty.copy(), np.full(periods + 1, np.nan), float("nan"))
    solution = result.x
    grid = solution[grid_offset:charge_offset]
    return DispatchResult(
        True,
        result.message,
        grid,
        solution[charge_offset:discharge_offset],
        solution[discharge_offset:pv_offset],
        solution[pv_offset:soc_offset],
        solution[soc_offset:down_offset],
        float(np.dot(prices, grid)),
    )


def run_rolling_day(
    load_kwh: np.ndarray | Sequence[float],
    actual_pv_kwh: np.ndarray | Sequence[float],
    hourly_forecast_kw: np.ndarray,
    price: np.ndarray | Sequence[float],
    initial_soc: float,
    terminal_soc: float,
    config: DispatchConfig = DispatchConfig(),
) -> RollingDayResult:
    load = np.asarray(load_kwh, dtype=float).reshape(-1)
    actual_pv = np.asarray(actual_pv_kwh, dtype=float).reshape(-1)
    prices = np.asarray(price, dtype=float).reshape(-1)
    forecasts = np.asarray(hourly_forecast_kw, dtype=float)
    if load.shape != (144,) or actual_pv.shape != (144,) or prices.shape != (144,):
        raise ValueError("rolling day requires exactly 144 ten-minute intervals")
    if forecasts.shape != (4, 24):
        raise ValueError("hourly_forecast_kw must have shape (4, 24)")

    midnight_pv_kwh = kw_to_interval_kwh(expand_hourly_forecast(forecasts[0]), config.interval_hours)
    plan = solve_dispatch(
        load,
        midnight_pv_kwh,
        prices,
        initial_soc=initial_soc,
        terminal_soc=terminal_soc,
        config=config,
    )
    if not plan.success:
        raise RuntimeError(f"0:00 plan optimization failed: {plan.message}")

    final_grid = plan.grid.copy()
    charge = plan.charge.copy()
    discharge = plan.discharge.copy()
    soc = plan.soc.copy()

    for release_index, release_hour in enumerate((6, 12, 18), start=1):
        start = release_hour * 6
        remaining_hours = 24 - release_hour
        revised_pv = kw_to_interval_kwh(
            expand_hourly_forecast(forecasts[release_index, :remaining_hours]),
            config.interval_hours,
        )
        revised = _solve_adjusted_dispatch(
            load[start:],
            revised_pv,
            prices[start:],
            plan.grid[start:],
            initial_soc=float(soc[start]),
            terminal_soc=terminal_soc,
            config=config,
        )
        if not revised.success:
            raise RuntimeError(f"{release_hour}:00 adjustment optimization failed: {revised.message}")
        final_grid[start:] = revised.grid
        charge[start:] = revised.charge
        discharge[start:] = revised.discharge
        soc[start:] = revised.soc

    required_pv = load + charge - final_grid - discharge
    if float(required_pv.min()) < -1e-5:
        raise RuntimeError("revised grid and storage schedule oversupply load even with zero PV")
    required_pv = np.maximum(required_pv, 0.0)
    pv_used_actual = np.minimum(actual_pv, required_pv)
    emergency = np.maximum(required_pv - actual_pv, 0.0)

    settlement = adjustment_settlement(plan.grid, final_grid, prices)
    emergency_cost = float(np.dot(5.0 * prices, emergency))
    total_cost = plan.purchase_cost + settlement.penalty_cost + settlement.incremental_cost + emergency_cost
    return RollingDayResult(
        plan=plan,
        final_grid=final_grid,
        charge=charge,
        discharge=discharge,
        pv_used_actual=pv_used_actual,
        emergency=emergency,
        soc=soc,
        adjustment=settlement,
        emergency_cost=emergency_cost,
        total_cost=total_cost,
    )


def solve_deterministic_year(
    data: InputData,
    price_matrix: np.ndarray,
    config: DispatchConfig = DispatchConfig(),
) -> DeterministicYearResult:
    """联合求解全年确定性调度，SOC 跨日连续且年初年末均为 6000 kWh。"""
    prices = np.asarray(price_matrix, dtype=float)
    if prices.shape != (365, 144):
        raise ValueError("price_matrix must have shape (365, 144)")
    load = kw_to_interval_kwh(data.load_kw.reshape(-1), config.interval_hours)
    pv = kw_to_interval_kwh(data.actual_pv_kw.reshape(-1), config.interval_hours)
    solved = solve_dispatch(load, pv, prices.reshape(-1), 6000.0, 6000.0, config)
    if not solved.success:
        raise RuntimeError(f"annual deterministic optimization failed: {solved.message}")
    grid = solved.grid.reshape(365, 144)
    charge = solved.charge.reshape(365, 144)
    discharge = solved.discharge.reshape(365, 144)
    pv_used = solved.pv_used.reshape(365, 144)
    soc = np.stack([solved.soc[day * 144 : (day + 1) * 144 + 1] for day in range(365)])
    daily_cost = np.sum(prices * grid, axis=1)
    return DeterministicYearResult(
        grid=grid,
        charge=charge,
        discharge=discharge,
        pv_used=pv_used,
        emergency=np.zeros_like(grid),
        soc=soc,
        daily_purchase_cost=daily_cost,
        total_cost=float(daily_cost.sum()),
    )


def solve_rolling_year(
    data: InputData,
    price_matrix: np.ndarray,
    config: DispatchConfig = DispatchConfig(),
    progress: bool = False,
) -> RollingYearResult:
    """逐日执行 0/6/12/18 点滚动优化；每天采用 SOC 周期边界避免末端透支。"""
    prices = np.asarray(price_matrix, dtype=float)
    if prices.shape != (365, 144):
        raise ValueError("price_matrix must have shape (365, 144)")
    arrays = [np.zeros((365, 144), dtype=float) for _ in range(6)]
    plan_grid, adjusted_grid, charge, discharge, pv_used, emergency = arrays
    soc = np.zeros((365, 145), dtype=float)
    daily_plan_cost = np.zeros(365)
    daily_adjustment_cost = np.zeros(365)
    daily_emergency_cost = np.zeros(365)
    downward = np.zeros(365)
    upward = np.zeros(365)

    initial_soc = 6000.0
    for day in range(365):
        result = run_rolling_day(
            kw_to_interval_kwh(data.load_kw[day], config.interval_hours),
            kw_to_interval_kwh(data.actual_pv_kw[day], config.interval_hours),
            data.forecast_pv_kw[day],
            prices[day],
            initial_soc=initial_soc,
            terminal_soc=initial_soc,
            config=config,
        )
        plan_grid[day] = result.plan.grid
        adjusted_grid[day] = result.final_grid
        charge[day] = result.charge
        discharge[day] = result.discharge
        pv_used[day] = result.pv_used_actual
        emergency[day] = result.emergency
        soc[day] = result.soc
        daily_plan_cost[day] = result.plan.purchase_cost
        daily_adjustment_cost[day] = result.adjustment.penalty_cost + result.adjustment.incremental_cost
        daily_emergency_cost[day] = result.emergency_cost
        downward[day] = result.adjustment.downward_kwh
        upward[day] = result.adjustment.upward_kwh
        initial_soc = float(result.soc[-1])
        if progress and (day + 1) % 30 == 0:
            print(f"rolling progress: {day + 1}/365 days", flush=True)

    return RollingYearResult(
        plan_grid, adjusted_grid, charge, discharge, pv_used, emergency, soc,
        daily_plan_cost, daily_adjustment_cost, daily_emergency_cost, downward, upward,
    )


def evaluate_midnight_only_year(
    data: InputData,
    price_matrix: np.ndarray,
    config: DispatchConfig = DispatchConfig(),
) -> dict[str, float]:
    """评估问题3只使用 0:00 预测且不在 6/12/18 点调整的对照方案。"""
    prices = np.asarray(price_matrix, dtype=float)
    if prices.shape != (365, 144):
        raise ValueError("price_matrix must have shape (365, 144)")
    plan_purchase = 0.0
    emergency_purchase = 0.0
    plan_cost = 0.0
    emergency_cost = 0.0
    for day in range(31, 365):
        load = kw_to_interval_kwh(data.load_kw[day], config.interval_hours)
        actual_pv = kw_to_interval_kwh(data.actual_pv_kw[day], config.interval_hours)
        forecast_pv = kw_to_interval_kwh(expand_hourly_forecast(data.forecast_pv_kw[day, 0]), config.interval_hours)
        plan = solve_dispatch(load, forecast_pv, prices[day], 6000.0, 6000.0, config)
        if not plan.success:
            raise RuntimeError(f"midnight-only baseline failed on day {day}: {plan.message}")
        required_pv = np.maximum(load + plan.charge - plan.grid - plan.discharge, 0.0)
        emergency = np.maximum(required_pv - actual_pv, 0.0)
        plan_purchase += float(plan.grid.sum())
        emergency_purchase += float(emergency.sum())
        plan_cost += plan.purchase_cost
        emergency_cost += float(np.dot(5.0 * prices[day], emergency))
    return {
        "plan_purchase_kwh": plan_purchase,
        "emergency_purchase_kwh": emergency_purchase,
        "total_physical_grid_purchase_kwh": plan_purchase + emergency_purchase,
        "plan_cost_yuan": plan_cost,
        "emergency_cost_yuan": emergency_cost,
        "total_cost_yuan": plan_cost + emergency_cost,
    }


def write_problem1_workbook(root: str | Path, output_path: str | Path, result: DispatchResult) -> Path:
    """复制并填写 result1.xlsx 模板。"""
    import shutil
    from openpyxl import load_workbook

    root_path = Path(root).resolve()
    destination = Path(output_path).resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    template = root_path / "附件" / "附件5" / "result1.xlsx"
    shutil.copy2(template, destination)
    workbook = load_workbook(destination)

    plan_sheet = workbook["计划购电量"]
    for row, (label, value) in enumerate(zip(build_time_labels(), result.grid), start=2):
        plan_sheet.cell(row, 1, label)
        plan_sheet.cell(row, 2, float(value))

    storage_sheet = workbook["充放电量"]
    for block in range(6):
        start = block * 24
        end = start + 24
        storage_sheet.cell(block + 2, 2, float(result.charge[start:end].sum()))
        storage_sheet.cell(block + 2, 3, float(result.discharge[start:end].sum()))
    storage_sheet.cell(2, 5, float(result.soc[0]))
    storage_sheet.cell(3, 5, float(result.soc[-1]))

    workbook.save(destination)
    return destination


def _copy_row_style(sheet, source_row: int, target_row: int, max_column: int) -> None:
    from copy import copy

    for column in range(1, max_column + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        target._style = copy(source._style)
        if source.has_style:
            target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
    source_dimension = sheet.row_dimensions[source_row]
    target_dimension = sheet.row_dimensions[target_row]
    target_dimension.height = source_dimension.height
    target_dimension.hidden = source_dimension.hidden


def _write_plan_sheet(sheet, values: np.ndarray, daily_cost: np.ndarray) -> None:
    for column, label in enumerate(build_time_labels(), start=2):
        sheet.cell(1, column, label)
    selected = slice(31, 365)
    for output_row, day in enumerate(range(selected.start, selected.stop), start=2):
        for slot, value in enumerate(values[day], start=2):
            sheet.cell(output_row, slot, float(value))
        sheet.cell(output_row, 146, float(values[day].sum()))
        sheet.cell(output_row, 147, float(daily_cost[day]))


def _write_storage_sheet(sheet, charge: np.ndarray, discharge: np.ndarray, soc: np.ndarray) -> None:
    from datetime import datetime, timedelta

    periods = ("0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00")
    day_count = 334
    target_last_row = 1 + 6 * day_count
    for day_output in range(day_count):
        day = day_output + 31
        row0 = 2 + 6 * day_output
        for block in range(6):
            row = row0 + block
            if row > 7:
                _copy_row_style(sheet, 2 + block, row, 6)
            for column in range(1, 7):
                sheet.cell(row, column).value = None
            if block == 0:
                sheet.cell(row, 1, datetime(2025, 2, 1) + timedelta(days=day_output))
            sheet.cell(row, 2, periods[block])
            start = 24 * block
            end = start + 24
            sheet.cell(row, 3, float(charge[day, start:end].sum()))
            sheet.cell(row, 4, float(discharge[day, start:end].sum()))
            if block == 0:
                sheet.cell(row, 5, datetime.strptime("0:00", "%H:%M").time())
                sheet.cell(row, 6, float(soc[day, 0]))
            elif block == 1:
                sheet.cell(row, 5, "24:00")
                sheet.cell(row, 6, float(soc[day, -1]))
    if sheet.max_row > target_last_row:
        sheet.delete_rows(target_last_row + 1, sheet.max_row - target_last_row)


def _write_emergency_sheet(sheet, emergency: np.ndarray) -> None:
    from datetime import datetime, timedelta

    output_row = 2
    for day_output in range(334):
        day = day_output + 31
        groups = group_emergency_intervals(emergency[day])
        rows_for_day = max(3, len(groups))
        for within_day in range(rows_for_day):
            if within_day == 0:
                source_row = 2
            elif within_day == rows_for_day - 1:
                source_row = 4
            else:
                source_row = 3
            if output_row > 4:
                _copy_row_style(sheet, source_row, output_row, 3)
            for column in range(1, 4):
                sheet.cell(output_row, column).value = None
            if within_day == 0:
                sheet.cell(output_row, 1, datetime(2025, 2, 1) + timedelta(days=day_output))
            if within_day < len(groups):
                label, amount = groups[within_day]
                sheet.cell(output_row, 2, label)
                sheet.cell(output_row, 3, float(amount))
            output_row += 1
    if sheet.max_row >= output_row:
        sheet.delete_rows(output_row, sheet.max_row - output_row + 1)


def write_deterministic_workbook(
    root: str | Path,
    output_path: str | Path,
    template_name: str,
    result: DeterministicYearResult,
    price_matrix: np.ndarray,
) -> Path:
    import shutil
    from openpyxl import load_workbook

    root_path = Path(root).resolve()
    destination = Path(output_path).resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(root_path / "附件" / "附件5" / template_name, destination)
    workbook = load_workbook(destination)
    _write_plan_sheet(workbook["计划购电量"], result.grid, result.daily_purchase_cost)
    _write_storage_sheet(workbook["充放电量"], result.charge, result.discharge, result.soc)
    _write_emergency_sheet(workbook["紧急购电量"], result.emergency)
    workbook.save(destination)
    return destination


def write_rolling_workbook(
    root: str | Path,
    output_path: str | Path,
    template_name: str,
    result: RollingYearResult,
    price_matrix: np.ndarray,
) -> Path:
    """复制并填写 result3.xlsx 或 result4-3.xlsx 模板。"""
    import shutil
    from openpyxl import load_workbook

    root_path = Path(root).resolve()
    destination = Path(output_path).resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(root_path / "附件" / "附件5" / template_name, destination)
    workbook = load_workbook(destination)
    _write_plan_sheet(
        workbook["计划购电量"],
        result.plan_grid,
        result.daily_plan_cost,
    )
    _write_plan_sheet(
        workbook["调整购电量"],
        result.adjusted_grid,
        result.daily_plan_cost + result.daily_adjustment_cost,
    )
    _write_storage_sheet(workbook["充放电量"], result.charge, result.discharge, result.soc)
    _write_emergency_sheet(workbook["紧急购电量"], result.emergency)
    workbook.save(destination)
    return destination


def audit_dispatch(
    grid: np.ndarray,
    charge: np.ndarray,
    discharge: np.ndarray,
    pv_used: np.ndarray,
    emergency: np.ndarray,
    load_kwh: np.ndarray,
    soc: np.ndarray,
    config: DispatchConfig = DispatchConfig(),
) -> dict[str, float | bool]:
    """独立复核功率平衡、SOC 方程、上下界、功率上限和互斥性。"""
    grid = np.asarray(grid, dtype=float)
    charge = np.asarray(charge, dtype=float)
    discharge = np.asarray(discharge, dtype=float)
    pv_used = np.asarray(pv_used, dtype=float)
    emergency = np.asarray(emergency, dtype=float)
    load = np.asarray(load_kwh, dtype=float)
    storage = np.asarray(soc, dtype=float)
    expected_shape = grid.shape
    if not all(array.shape == expected_shape for array in (charge, discharge, pv_used, emergency, load)):
        raise ValueError("dispatch arrays must have the same shape")
    if storage.shape[:-1] != expected_shape[:-1] or storage.shape[-1] != expected_shape[-1] + 1:
        raise ValueError("SOC must have one more point than interval arrays")

    balance = grid + pv_used + discharge + emergency - load - charge
    soc_error = storage[..., 1:] - storage[..., :-1] - config.efficiency * charge + discharge / config.efficiency
    tolerance = 1e-5
    return {
        "max_power_balance_error_kwh": float(np.max(np.abs(balance))),
        "max_soc_dynamics_error_kwh": float(np.max(np.abs(soc_error))),
        "soc_min_kwh": float(storage.min()),
        "soc_max_kwh": float(storage.max()),
        "soc_bounds_ok": bool(storage.min() >= config.soc_min_kwh - tolerance and storage.max() <= config.soc_max_kwh + tolerance),
        "max_charge_interval_kwh": float(charge.max()),
        "max_discharge_interval_kwh": float(discharge.max()),
        "power_limits_ok": bool(charge.max() <= config.interval_limit_kwh + tolerance and discharge.max() <= config.interval_limit_kwh + tolerance),
        "max_simultaneous_charge_discharge_kwh": float(np.max(np.minimum(charge, discharge))),
        "no_simultaneous_charge_discharge": bool(np.max(np.minimum(charge, discharge)) <= tolerance),
        "initial_soc_kwh": float(storage.reshape(-1, storage.shape[-1])[0, 0]),
        "final_soc_kwh": float(storage.reshape(-1, storage.shape[-1])[-1, -1]),
    }


def _deterministic_metrics(
    result: DeterministicYearResult,
    actual_pv_kwh: np.ndarray,
    start_day: int = 31,
) -> dict[str, float]:
    selected = slice(start_day, 365)
    return {
        "plan_purchase_kwh": float(result.grid[selected].sum()),
        "adjusted_purchase_kwh": float(result.grid[selected].sum()),
        "emergency_purchase_kwh": float(result.emergency[selected].sum()),
        "total_physical_grid_purchase_kwh": float((result.grid[selected] + result.emergency[selected]).sum()),
        "total_cost_yuan": float(result.daily_purchase_cost[selected].sum()),
        "charge_kwh": float(result.charge[selected].sum()),
        "discharge_kwh": float(result.discharge[selected].sum()),
        "pv_consumed_kwh": float(result.pv_used[selected].sum()),
        "pv_curtailed_kwh": float(actual_pv_kwh[selected].sum() - result.pv_used[selected].sum()),
    }


def _rolling_metrics(
    result: RollingYearResult,
    actual_pv_kwh: np.ndarray,
    start_day: int = 31,
) -> dict[str, float]:
    selected = slice(start_day, 365)
    return {
        "plan_purchase_kwh": float(result.plan_grid[selected].sum()),
        "adjusted_purchase_kwh": float(result.adjusted_grid[selected].sum()),
        "downward_adjustment_kwh": float(result.downward_kwh[selected].sum()),
        "upward_adjustment_kwh": float(result.upward_kwh[selected].sum()),
        "emergency_purchase_kwh": float(result.emergency[selected].sum()),
        "total_physical_grid_purchase_kwh": float((result.adjusted_grid[selected] + result.emergency[selected]).sum()),
        "plan_cost_yuan": float(result.daily_plan_cost[selected].sum()),
        "adjustment_cost_yuan": float(result.daily_adjustment_cost[selected].sum()),
        "emergency_cost_yuan": float(result.daily_emergency_cost[selected].sum()),
        "total_cost_yuan": float(result.daily_total_cost[selected].sum()),
        "charge_kwh": float(result.charge[selected].sum()),
        "discharge_kwh": float(result.discharge[selected].sum()),
        "pv_consumed_kwh": float(result.pv_used_actual[selected].sum()),
        "pv_curtailed_kwh": float(actual_pv_kwh[selected].sum() - result.pv_used_actual[selected].sum()),
    }


def build_summary(
    data: InputData,
    problem1: DispatchResult,
    problem2: object,
    problem3: RollingYearResult,
    problem4_2: DeterministicYearResult,
    problem4_3: RollingYearResult,
    problem3_midnight_only: dict[str, float],
    problem4_3_midnight_only: dict[str, float],
) -> dict:
    from problem2_optimization import Problem2Data, problem2_metrics

    actual_pv_kwh = kw_to_interval_kwh(data.actual_pv_kw)
    labels = build_time_labels()
    specified = ("10:00-10:10", "12:00-12:10", "14:00-14:10", "16:00-16:10", "18:00-18:10", "20:00-20:10")
    key_dates = ("2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21")
    date_to_index = {str(value): index for index, value in enumerate(data.dates)}
    results = {
        "assumptions": {
            "interval_minutes": 10,
            "power_to_energy_factor_hours": 1.0 / 6.0,
            "charge_efficiency": 0.9,
            "discharge_efficiency": 0.9,
            "soc_bounds_kwh": [1200.0, 10800.0],
            "power_limit_kw": 5000.0,
            "problem2_forecast": "causal same-slot lag ensemble plus rolling residual quantile; alpha selected only from prior-day validation",
            "problem3_forecast_conversion": "zero-order hold: each hourly forecast is repeated for six 10-minute intervals",
            "time_alignment": "attachment timestamps are interval ends: 0:10 maps to 0:00-0:10 and 0:00+1 maps to 23:50-24:00",
            "rolling_times": ["0:00", "6:00", "12:00", "18:00"],
            "rolling_terminal_policy": "daily terminal SOC equals that day's initial SOC",
            "result_period": "2025-02-01 through 2025-12-31",
        },
        "problem1": {
            "total_purchase_kwh": float(problem1.grid.sum()),
            "total_cost_yuan": float(problem1.purchase_cost),
            "charge_kwh": float(problem1.charge.sum()),
            "discharge_kwh": float(problem1.discharge.sum()),
            "pv_consumed_kwh": float(problem1.pv_used.sum()),
            "pv_curtailed_kwh": float(kw_to_interval_kwh(data.problem1_pv_kw).sum() - problem1.pv_used.sum()),
            "specified_interval_purchase_kwh": {label: float(problem1.grid[labels.index(label)]) for label in specified},
        },
        "problem2": problem2_metrics(
            problem2,
            Problem2Data(data.static_price, data.dates, data.load_kw, data.actual_pv_kw),
        ),
        "problem3": _rolling_metrics(problem3, actual_pv_kwh),
        "problem4_2": _deterministic_metrics(problem4_2, actual_pv_kwh),
        "problem4_3": _rolling_metrics(problem4_3, actual_pv_kwh),
    }
    for name, result in (("problem2", problem2), ("problem3", problem3), ("problem4_2", problem4_2), ("problem4_3", problem4_3)):
        emergency = result.emergency
        results[name]["key_date_emergency_kwh"] = {
            date: float(emergency[date_to_index[date]].sum()) for date in key_dates
        }
    for name, baseline in (("problem3", problem3_midnight_only), ("problem4_3", problem4_3_midnight_only)):
        results[name]["midnight_only_baseline"] = baseline
        results[name]["rolling_forecast_benefit"] = {
            "emergency_reduction_kwh": baseline["emergency_purchase_kwh"] - results[name]["emergency_purchase_kwh"],
            "cost_reduction_yuan": baseline["total_cost_yuan"] - results[name]["total_cost_yuan"],
            "cost_reduction_percent": 100.0 * (baseline["total_cost_yuan"] - results[name]["total_cost_yuan"]) / baseline["total_cost_yuan"],
        }
    return results


def write_summary_files(output_dir: str | Path, summary: dict, validation: dict) -> None:
    import csv
    import json

    destination = Path(output_dir).resolve()
    destination.mkdir(parents=True, exist_ok=True)
    with (destination / "summary.json").open("w", encoding="utf-8") as handle:
        json.dump(summary, handle, ensure_ascii=False, indent=2)
    with (destination / "validation_report.json").open("w", encoding="utf-8") as handle:
        json.dump(validation, handle, ensure_ascii=False, indent=2)
    with (destination / "summary.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(("问题", "指标", "数值"))
        for question, metrics in summary.items():
            if question == "assumptions":
                continue
            for metric, value in metrics.items():
                if isinstance(value, dict):
                    for key, nested_value in value.items():
                        writer.writerow((question, f"{metric}.{key}", nested_value))
                else:
                    writer.writerow((question, metric, value))


def validate_workbook_contract(output_path: str | Path, template_path: str | Path) -> dict[str, object]:
    from openpyxl import load_workbook

    output = load_workbook(output_path, read_only=False, data_only=False)
    template = load_workbook(template_path, read_only=False, data_only=False)
    names_ok = output.sheetnames == template.sheetnames
    headers_ok = True
    header_styles_ok = True
    interval_labels_ok = True
    is_problem1 = output_path.__str__().lower().endswith("result1.xlsx")
    aligned_labels = build_time_labels()
    for sheet_name in template.sheetnames:
        template_sheet = template[sheet_name]
        output_sheet = output[sheet_name]
        max_column = template_sheet.max_column
        for column in range(1, max_column + 1):
            expected_header = template_sheet.cell(1, column).value
            if not is_problem1 and sheet_name in ("计划购电量", "调整购电量") and 2 <= column <= 145:
                expected_header = aligned_labels[column - 2]
            headers_ok = headers_ok and output_sheet.cell(1, column).value == expected_header
        header_styles_ok = header_styles_ok and all(
            output_sheet.cell(1, column)._style == template_sheet.cell(1, column)._style
            for column in range(1, max_column + 1)
        )

    if is_problem1:
        plan_sheet = output["计划购电量"]
        interval_labels_ok = all(plan_sheet.cell(row, 1).value == aligned_labels[row - 2] for row in range(2, 146))
    else:
        for sheet_name in ("计划购电量", "调整购电量"):
            if sheet_name in output.sheetnames:
                interval_labels_ok = interval_labels_ok and all(
                    output[sheet_name].cell(1, column).value == aligned_labels[column - 2]
                    for column in range(2, 146)
                )

    plan_sheets = [name for name in ("计划购电量", "调整购电量") if name in output.sheetnames]
    plan_complete = True
    for sheet_name in plan_sheets:
        sheet = output[sheet_name]
        if is_problem1:
            plan_complete = plan_complete and all(sheet.cell(row, 2).value is not None for row in range(2, 146))
        else:
            plan_complete = plan_complete and all(
                sheet.cell(row, column).value is not None
                for row in range(2, 336)
                for column in range(2, 148)
            )

    storage_dates = None
    emergency_dates = None
    if "充放电量" in output.sheetnames and not output_path.__str__().lower().endswith("result1.xlsx"):
        storage_dates = sum(output["充放电量"].cell(row, 1).value is not None for row in range(2, output["充放电量"].max_row + 1))
    if "紧急购电量" in output.sheetnames:
        emergency_dates = sum(output["紧急购电量"].cell(row, 1).value is not None for row in range(2, output["紧急购电量"].max_row + 1))
    result = {
        "sheet_names_ok": names_ok,
        "headers_ok": headers_ok,
        "interval_labels_ok": interval_labels_ok,
        "header_styles_ok": header_styles_ok,
        "plan_cells_complete": plan_complete,
        "storage_date_count": storage_dates,
        "emergency_date_count": emergency_dates,
        "complete": bool(names_ok and headers_ok and interval_labels_ok and header_styles_ok and plan_complete and (storage_dates in (None, 334)) and (emergency_dates in (None, 334))),
    }
    output.close()
    template.close()
    return result


def build_numeric_validation(
    data: InputData,
    problem1: DispatchResult,
    problem2: object,
    problem3: RollingYearResult,
    problem4_2: DeterministicYearResult,
    problem4_3: RollingYearResult,
) -> dict[str, dict[str, float | bool]]:
    from problem2_optimization import Problem2Data, audit_problem2

    load1 = kw_to_interval_kwh(data.problem1_load_kw)
    annual_load = kw_to_interval_kwh(data.load_kw)
    zero1 = np.zeros(144)
    validation = {
        "problem1": audit_dispatch(problem1.grid, problem1.charge, problem1.discharge, problem1.pv_used, zero1, load1, problem1.soc),
        "problem2": audit_problem2(
            problem2,
            Problem2Data(data.static_price, data.dates, data.load_kw, data.actual_pv_kw),
        ),
        "problem3": audit_dispatch(problem3.adjusted_grid, problem3.charge, problem3.discharge, problem3.pv_used_actual, problem3.emergency, annual_load, problem3.soc),
        "problem4_2": audit_dispatch(problem4_2.grid, problem4_2.charge, problem4_2.discharge, problem4_2.pv_used, problem4_2.emergency, annual_load, problem4_2.soc),
        "problem4_3": audit_dispatch(problem4_3.adjusted_grid, problem4_3.charge, problem4_3.discharge, problem4_3.pv_used_actual, problem4_3.emergency, annual_load, problem4_3.soc),
    }
    validation["problem1"]["boundary_soc_ok"] = bool(abs(problem1.soc[0] - 6000.0) <= 1e-5 and abs(problem1.soc[-1] - 6000.0) <= 1e-5)
    for name, result in (("problem4_2", problem4_2),):
        continuity_error = float(np.max(np.abs(result.soc[:-1, -1] - result.soc[1:, 0])))
        validation[name]["cross_day_soc_continuity_error_kwh"] = continuity_error
        validation[name]["boundary_soc_ok"] = bool(abs(result.soc[0, 0] - 6000.0) <= 1e-5 and abs(result.soc[-1, -1] - 6000.0) <= 1e-5 and continuity_error <= 1e-5)
    for name, result in (("problem3", problem3), ("problem4_3", problem4_3)):
        daily_error = float(np.max(np.abs(result.soc[:, 0] - result.soc[:, -1])))
        validation[name]["daily_boundary_soc_error_kwh"] = daily_error
        validation[name]["boundary_soc_ok"] = bool(abs(result.soc[0, 0] - 6000.0) <= 1e-5 and abs(result.soc[-1, -1] - 6000.0) <= 1e-5 and daily_error <= 1e-5)
    return validation


def main(argv: Sequence[str] | None = None) -> int:
    import argparse

    parser = argparse.ArgumentParser(description="Solve all four 2026 C-problem microgrid optimization tasks.")
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parent, help="目录中应包含 C题.pdf 和 附件/")
    parser.add_argument("--output-dir", type=Path, default=None, help="结果目录；默认写入 root")
    parser.add_argument("--quiet", action="store_true", help="不打印每 30 天一次的滚动进度")
    arguments = parser.parse_args(argv)
    root = arguments.root.resolve()
    output_dir = (arguments.output_dir or root).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    print("[1/8] 读取并校验附件 1-4", flush=True)
    data = load_input_data(root)
    static_prices = np.broadcast_to(data.static_price, (365, 144)).copy()

    print("[2/8] 求解问题1", flush=True)
    problem1 = solve_dispatch(
        kw_to_interval_kwh(data.problem1_load_kw),
        kw_to_interval_kwh(data.problem1_pv_kw),
        data.static_price,
        6000.0,
        6000.0,
    )
    if not problem1.success:
        raise RuntimeError(problem1.message)
    write_problem1_workbook(root, output_dir / "result1.xlsx", problem1)

    print("[3/8] 求解问题2（严格因果日前计划与风险修正回测）", flush=True)
    from problem2_optimization import Problem2Data, run_backtest, save_result2

    problem2_data = Problem2Data(data.static_price, data.dates, data.load_kw, data.actual_pv_kw)
    problem2 = run_backtest(problem2_data, progress=not arguments.quiet)
    save_result2(root, output_dir / "result2.xlsx", problem2)

    print("[4/8] 求解问题4-2（动态电价、全年联合优化）", flush=True)
    problem4_2 = solve_deterministic_year(data, data.dynamic_price)
    write_deterministic_workbook(root, output_dir / "result4-2.xlsx", "result4-2.xlsx", problem4_2, data.dynamic_price)

    print("[5/8] 求解问题3（固定电价、0/6/12/18点滚动优化）", flush=True)
    problem3 = solve_rolling_year(data, static_prices, progress=not arguments.quiet)
    write_rolling_workbook(root, output_dir / "result3.xlsx", "result3.xlsx", problem3, static_prices)

    print("[6/8] 求解问题4-3（动态电价、0/6/12/18点滚动优化）", flush=True)
    problem4_3 = solve_rolling_year(data, data.dynamic_price, progress=not arguments.quiet)
    write_rolling_workbook(root, output_dir / "result4-3.xlsx", "result4-3.xlsx", problem4_3, data.dynamic_price)

    print("[7/9] 执行数值和工作簿完整性检查", flush=True)
    numeric_validation = build_numeric_validation(data, problem1, problem2, problem3, problem4_2, problem4_3)
    workbook_validation = {}
    for filename in ("result1.xlsx", "result2.xlsx", "result3.xlsx", "result4-2.xlsx", "result4-3.xlsx"):
        workbook_validation[filename] = validate_workbook_contract(
            output_dir / filename,
            root / "附件" / "附件5" / filename,
        )
    numeric_ok = all(
        bool(item["soc_bounds_ok"])
        and bool(item["power_limits_ok"])
        and bool(item["no_simultaneous_charge_discharge"])
        and bool(item["boundary_soc_ok"])
        and float(item["max_power_balance_error_kwh"]) <= 1e-5
        and float(item["max_soc_dynamics_error_kwh"]) <= 1e-5
        for item in numeric_validation.values()
    )
    workbook_ok = all(bool(item["complete"]) for item in workbook_validation.values())
    validation = {
        "numeric": numeric_validation,
        "workbooks": workbook_validation,
        "all_checks_passed": bool(numeric_ok and workbook_ok),
    }
    print("[8/9] 计算仅使用0:00预测的对照方案", flush=True)
    problem3_midnight_only = evaluate_midnight_only_year(data, static_prices)
    problem4_3_midnight_only = evaluate_midnight_only_year(data, data.dynamic_price)
    summary = build_summary(
        data,
        problem1,
        problem2,
        problem3,
        problem4_2,
        problem4_3,
        problem3_midnight_only,
        problem4_3_midnight_only,
    )
    write_summary_files(output_dir, summary, validation)
    if not validation["all_checks_passed"]:
        raise RuntimeError("validation failed; inspect validation_report.json")

    print("[9/9] 全部结果和审计文件已生成", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
